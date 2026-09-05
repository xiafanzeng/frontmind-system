#!/usr/bin/env python3
"""Import an exactly-one-question Dashboard citation workbook without reranking."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
import unicodedata
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared"))
from schema_validation import validate_root  # noqa: E402
from xlsx_safety import validate_xlsx  # noqa: E402

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - dependency preflight owns this path
    raise RuntimeError("openpyxl is required; run the workflow dependency preflight") from exc


REQUIRED_SHEETS = ("内容统计", "详细表格", "字段说明")
CONTENT_HEADERS = (
    "引用排行", "内容状态", "内容标题", "内容URL", "媒体名称", "媒体url", "引用日期", "引用次数", "引用占比",
)
DETAIL_HEADERS = ("引用日期", "模型", "监控问题", "内容标题", "内容链接", "媒体名称", "媒体url")
GUIDE_HEADERS = ("字段名称", "字段释义")
REQUIRED_GUIDE_FIELDS = {
    "引用排行", "引用占比", "内容标题", "内容URL", "引用次数", "内容状态",
}
TRACKING_QUERY_KEYS = {
    "utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content",
    "gclid", "fbclid", "spm", "from", "source",
}
ENTRY_CATEGORIES = {"industry_ranking", "competitor_comparison", "reputation", "product_scenario"}
PRODUCT_SUBINTENTS = {
    "offering_definition", "feature_mechanism", "scenario_fit", "delivery_usage", "support_boundary", "price_selection",
}


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalized_question(value: Any) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text(value))).strip()


def normalized_definition(value: Any) -> str:
    """Normalize guide prose while retaining only characters with semantic value."""
    normalized = unicodedata.normalize("NFKC", text(value)).casefold()
    return "".join(
        character for character in normalized
        if not unicodedata.category(character).startswith(("P", "S", "Z"))
    )


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def relative_inside(root: Path, path: Path, label: str) -> str:
    root = root.resolve()
    resolved = path.resolve(strict=False)
    try:
        return resolved.relative_to(root).as_posix()
    except ValueError as exc:
        raise ValueError(f"{label} must be inside --package-root") from exc


def atomic_json(path: Path, value: dict[str, Any]) -> None:
    if path.exists() or path.is_symlink():
        raise FileExistsError(f"refusing to overwrite E2 source artifact: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def header_map(sheet, expected: tuple[str, ...]) -> dict[str, int]:
    headers = [text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if len(headers) != len(set(headers)):
        raise ValueError(f"{sheet.title} contains duplicate headers")
    missing = [name for name in expected if name not in headers]
    if missing:
        raise ValueError(f"{sheet.title} is missing required Dashboard columns: {missing}")
    return {name: headers.index(name) for name in expected}


def row_values(row: tuple[Any, ...]) -> list[Any]:
    return [cell.value for cell in row]


def canonical_url(value: Any) -> str:
    raw = text(value)
    if not raw:
        raise ValueError("Dashboard content URL is empty")
    parsed = urlsplit(raw)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
        # Unsafe/non-HTTP rows remain in the raw Top20 and are finalized as
        # unsafe_url during page analysis; never backfill from rank 21.
        return raw
    hostname = (parsed.hostname or "").lower()
    port = parsed.port
    netloc = hostname
    if port and not ((parsed.scheme.lower() == "http" and port == 80) or (parsed.scheme.lower() == "https" and port == 443)):
        netloc = f"{hostname}:{port}"
    query = urlencode(
        sorted((key, value) for key, value in parse_qsl(parsed.query, keep_blank_values=True)
               if key.lower() not in TRACKING_QUERY_KEYS),
        doseq=True,
    )
    return urlunsplit((parsed.scheme.lower(), netloc, parsed.path or "/", query, ""))


def integer(value: Any, label: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{label} must be an integer")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be an integer") from exc
    if isinstance(value, float) and value != number:
        raise ValueError(f"{label} must be an integer")
    return number


def share(value: Any, label: str) -> float:
    raw = text(value)
    try:
        if raw.endswith("%"):
            result = float(raw[:-1]) / 100
        else:
            result = float(value)
            if result > 1:
                result /= 100
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be a percentage or 0-1 number") from exc
    if not 0 <= result <= 1:
        raise ValueError(f"{label} must be between 0 and 1")
    return round(result, 12)


def iso_midnight(value: str) -> str:
    parsed = datetime.strptime(value, "%Y-%m-%d").replace(tzinfo=timezone.utc)
    return parsed.isoformat()


def date_range(value: Any) -> dict[str, str | None]:
    if value is None or text(value) == "":
        return {"from": None, "to": None}
    if isinstance(value, datetime):
        stamp = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return {"from": stamp.isoformat(), "to": stamp.isoformat()}
    if isinstance(value, date):
        stamp = datetime.combine(value, time.min, tzinfo=timezone.utc).isoformat()
        return {"from": stamp, "to": stamp}
    dates = re.findall(r"\d{4}-\d{2}-\d{2}", text(value))
    if not dates:
        raise ValueError(f"引用日期 must contain YYYY-MM-DD values: {value!r}")
    return {"from": iso_midnight(dates[0]), "to": iso_midnight(dates[-1])}


def media_domain(value: Any) -> str | None:
    raw = text(value)
    if not raw:
        return None
    parsed = urlsplit(raw if "://" in raw else f"https://{raw}")
    return (parsed.hostname or "").lower() or None


def effective_range(sheet) -> str:
    max_row = sheet.max_row or 0
    max_column = sheet.max_column or 0
    if not max_row or not max_column:
        # Some valid OOXML producers omit worksheet dimension metadata.
        # openpyxl read-only mode then reports None, so deterministically scan
        # values rather than rejecting an otherwise valid Dashboard export.
        max_row = 0
        max_column = 0
        for row_number, row in enumerate(sheet.iter_rows(), 1):
            populated_columns = [
                column_number for column_number, cell in enumerate(row, 1)
                if cell.value not in {None, ""}
            ]
            if populated_columns:
                max_row = row_number
                max_column = max(max_column, max(populated_columns))
    max_row = max(1, max_row)
    max_column = max(1, max_column)
    return f"A1:{get_column_letter(max_column)}{max_row}"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--content-job", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--package-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    package_root = args.package_root.resolve()
    if args.content_job.is_symlink():
        raise ValueError("canonical Content Job must be a regular non-symlink file")
    content_job_path = args.content_job.resolve()
    if not package_root.is_dir() or not content_job_path.is_file():
        raise FileNotFoundError("package root or regular canonical Content Job is missing")
    relative_inside(package_root, content_job_path, "--content-job")
    content_job = json.loads(content_job_path.read_text(encoding="utf-8"))
    if not isinstance(content_job, dict):
        raise ValueError("--content-job must contain a JSON object")
    validate_root(content_job, "content_job.schema.json", "E2 canonical Content Job")
    entry_category = content_job.get("entry_category")
    if entry_category not in ENTRY_CATEGORIES or entry_category == "foundation_start":
        raise ValueError("Dashboard E2 adapter only accepts the four ordinary Content Job entries")
    task = content_job.get("task") if isinstance(content_job.get("task"), dict) else {}
    question = normalized_question(task.get("question_text"))
    product_scenario_subintent = task.get("product_scenario_subintent")
    if task.get("question_text") != question:
        raise ValueError(
            "canonical Content Job question_text is not NFKC/whitespace normalized; rerun the E1 normalizer"
        )
    if content_job.get("source_workbook_path") != "01_monitoring/dashboard_source_workbook.xlsx":
        raise ValueError(
            "ordinary Content Job source_workbook_path must be "
            "01_monitoring/dashboard_source_workbook.xlsx"
        )
    if not 4 <= len(question) <= 300:
        raise ValueError("Content Job question_text must normalize to 4-300 characters")
    if entry_category == "product_scenario" and product_scenario_subintent not in PRODUCT_SUBINTENTS:
        raise ValueError("product_scenario Content Job requires one canonical product_scenario_subintent")
    if entry_category != "product_scenario" and product_scenario_subintent is not None:
        raise ValueError("only product_scenario Content Job may set product_scenario_subintent")
    if args.workbook.is_symlink():
        raise ValueError("staged Dashboard workbook must be a regular non-symlink file")
    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError("package root or regular staged Dashboard workbook is missing")
    validate_xlsx(workbook_path)
    workbook_relative = relative_inside(package_root, workbook_path, "--workbook")
    if content_job.get("source_workbook_path") != workbook_relative:
        raise ValueError(
            "--workbook must be the exact job-relative source_workbook_path bound by the canonical Content Job"
        )
    output_relative = relative_inside(package_root, args.output, "--output")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        missing_sheets = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
        if missing_sheets:
            raise ValueError(f"Dashboard workbook is missing required sheets: {missing_sheets}")
        content_sheet, detail_sheet, guide_sheet = (workbook[name] for name in REQUIRED_SHEETS)
        content_columns = header_map(content_sheet, CONTENT_HEADERS)
        detail_columns = header_map(detail_sheet, DETAIL_HEADERS)
        guide_columns = header_map(guide_sheet, GUIDE_HEADERS)

        guide_definitions = {
            text(values[guide_columns["字段名称"]]): text(values[guide_columns["字段释义"]])
            for values in (row_values(row) for row in guide_sheet.iter_rows(min_row=2))
            if any(value not in {None, ""} for value in values)
        }
        guide_fields = set(guide_definitions)
        missing_guide = sorted(REQUIRED_GUIDE_FIELDS - guide_fields)
        if missing_guide:
            raise ValueError(f"字段说明 does not define required Dashboard fields: {missing_guide}")
        rank_definition = normalized_definition(guide_definitions.get("引用排行"))
        if not (
            re.search(r"按.{0,8}(?:被引|引用)次数.{0,8}排序", rank_definition)
            or re.search(r"(?:被引|引用)次数.{0,8}(?:降序|由高到低|从高到低)", rank_definition)
        ):
            raise ValueError("字段说明.引用排行 must explicitly define ranking by citation count")
        share_definition = normalized_definition(guide_definitions.get("引用占比"))
        if not (
            ("信源引用次数" in share_definition or "该信源引用次数" in share_definition)
            and "总引用次数" in share_definition
            and ("百分比" in share_definition or "占比" in share_definition)
            and "影响力" in share_definition
        ):
            raise ValueError(
                "字段说明.引用占比 must define the source citation count as a percentage "
                "of total citations and explain that it measures source influence"
            )

        detail_rows: list[tuple[int, list[Any]]] = []
        observed_questions: dict[str, int] = {}
        refs_by_original: dict[str, list[int]] = {}
        refs_by_canonical: dict[str, list[int]] = {}
        for row_number, row in enumerate(detail_sheet.iter_rows(min_row=2), 2):
            values = row_values(row)
            if not any(value not in {None, ""} for value in values):
                continue
            row_question = normalized_question(values[detail_columns["监控问题"]])
            if not row_question:
                raise ValueError(f"详细表格 row {row_number} has no 监控问题")
            observed_questions[row_question] = observed_questions.get(row_question, 0) + 1
            raw_url = text(values[detail_columns["内容链接"]])
            if not raw_url:
                raise ValueError(f"详细表格 row {row_number} has no 内容链接")
            normalized_url = canonical_url(raw_url)
            refs_by_original.setdefault(raw_url, []).append(row_number)
            refs_by_canonical.setdefault(normalized_url, []).append(row_number)
            detail_rows.append((row_number, values))
        if set(observed_questions) != {question}:
            preview = ", ".join(f"{key!r}({count})" for key, count in sorted(observed_questions.items()))
            raise ValueError(
                "Dashboard workbook must be prefiltered to exactly the Content Job question; "
                f"expected {question!r}, found {preview or 'no questions'}"
            )

        ranked: dict[int, dict[str, Any]] = {}
        all_export_ranks: set[int] = set()
        all_rank_counts: dict[int, int] = {}
        for row_number, row in enumerate(content_sheet.iter_rows(min_row=2), 2):
            values = row_values(row)
            if not any(value not in {None, ""} for value in values):
                continue
            rank = integer(values[content_columns["引用排行"]], f"内容统计 row {row_number} 引用排行")
            if rank < 1:
                raise ValueError(f"内容统计 row {row_number} 引用排行 must be positive")
            if rank in all_export_ranks:
                raise ValueError(f"内容统计 contains duplicate 引用排行 {rank}")
            all_export_ranks.add(rank)
            count = integer(values[content_columns["引用次数"]], f"内容统计 row {row_number} 引用次数")
            if count < 0:
                raise ValueError(f"内容统计 row {row_number} 引用次数 must be non-negative")
            all_rank_counts[rank] = count
            if rank > 20:
                continue
            raw_url = text(values[content_columns["内容URL"]])
            normalized_url = canonical_url(raw_url)
            refs = refs_by_original.get(raw_url) or refs_by_canonical.get(normalized_url) or []
            # The exported count field is non-negative by contract.  A row can
            # enter the cited Top20 only when its URL is also present in the
            # exact-question detail table, which independently implies >=1.
            if not refs:
                raise ValueError(
                    f"raw Top20 rank {rank} URL has no exact-question 详细表格 record"
                )
            if len(refs) != count:
                raise ValueError(
                    f"raw Top20 rank {rank} count mismatch: 内容统计={count}, exact-question 详细表格 rows={len(refs)}; "
                    "the workbook is not a consistent prefiltered single-question export"
                )
            title = text(values[content_columns["内容标题"]])
            if not title:
                raise ValueError(f"raw Top20 rank {rank} has no 内容标题")
            ranked[rank] = {
                "raw_rank": rank,
                "content_title": title,
                "original_url": raw_url,
                "canonical_url": normalized_url,
                "media_name": text(values[content_columns["媒体名称"]]) or None,
                "media_domain": media_domain(values[content_columns["媒体url"]]),
                "citation_date_range": date_range(values[content_columns["引用日期"]]),
                "citation_count": count,
                "citation_share": share(values[content_columns["引用占比"]], f"内容统计 row {row_number} 引用占比"),
                "detailed_row_refs": sorted(set(refs)),
            }
        if not ranked:
            raise ValueError("内容统计 contains no raw ranks 1-20")
        ordered_rank_counts = [all_rank_counts[rank] for rank in sorted(all_rank_counts)]
        if any(left < right for left, right in zip(ordered_rank_counts, ordered_rank_counts[1:])):
            raise ValueError("内容统计.引用排行 is not citation_count descending as declared")
        expected_contiguous = set(range(1, max(ranked) + 1))
        if set(ranked) != expected_contiguous:
            raise ValueError(
                f"raw Dashboard Top20 ranks must be contiguous 1..N without gaps; found={sorted(ranked)}"
            )
        exact_question_total = len(detail_rows)
        for rank, item in ranked.items():
            expected_share = item["citation_count"] / exact_question_total
            # Dashboard displays citation share at two decimal percentage
            # points, so half a display unit is the only accepted tolerance.
            if abs(item["citation_share"] - expected_share) > 0.000051:
                raise ValueError(
                    f"raw Top20 rank {rank} citation_share does not reconcile to the exact-question detail total: "
                    f"export={item['citation_share']}, expected={expected_share}"
                )
        pool = [ranked[rank] for rank in sorted(ranked)]
        pool_status = "complete_pool" if set(ranked) == set(range(1, 21)) else "partial_pool"
        payload = {
            "schema_version": "2.1.0",
            "artifact_type": "E2_dashboard_source_pool",
            "job_id": content_job.get("job_id"),
            "question": question,
            "entry_category": entry_category,
            "product_scenario_subintent": product_scenario_subintent,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "source_workbook_receipt": {
                "source_path": workbook_relative,
                "sha256": file_sha256(workbook_path),
                "sheet_names": list(workbook.sheetnames),
                "effective_ranges": {name: effective_range(workbook[name]) for name in REQUIRED_SHEETS},
                "normalized_question": question,
                "content_status_ignored": True,
                "ranking_basis": "dashboard_citation_rank",
                "citation_rank_definition": "citation_count_descending",
                "dashboard_total_citation_count": exact_question_total,
                "top20_citation_count": sum(item["citation_count"] for item in pool),
            },
            "pool_status": pool_status,
            "cited_content_pool": pool,
        }
    finally:
        workbook.close()

    if "内容状态" in json.dumps(payload, ensure_ascii=False):
        raise AssertionError("Dashboard 内容状态 leaked into the E2 source artifact")
    atomic_json(package_root / output_relative, payload)
    print(json.dumps({
        "status": "passed", "job_id": content_job.get("job_id"), "question": question,
        "pool_status": pool_status, "raw_pool_size": len(pool), "output": output_relative,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
