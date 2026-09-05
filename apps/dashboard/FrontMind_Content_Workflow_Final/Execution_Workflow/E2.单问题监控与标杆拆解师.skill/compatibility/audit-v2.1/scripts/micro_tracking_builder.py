#!/usr/bin/env python3
"""Build the answer-landscape monitoring input from two tabular exports.

This artifact preserves answers, cited entities and user-question gaps only.
Dashboard Top20 retrieval, P01-P16 page classification, structural benchmark
analysis and recommendation belong exclusively to the v2.1 E2 source chain.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared"))
from schema_validation import validate_root  # noqa: E402


ANSWER_ALIASES = {
    "answer_id": ("answer_id", "observation_id", "答案ID", "sourceRecordId", "source_record_id"),
    "question": ("question", "问题", "问题原文", "query"),
    "question_id": ("questionId", "question_id", "问题ID"),
    "platform": ("platform", "平台", "engine", "模型平台"),
    "model": ("model", "模型", "model_name"),
    "answer_no": ("answerNo", "answer_no", "答案序号"),
    "answer_text": ("answer", "答案", "回答", "response", "content"),
    "citation_count": ("citationCount", "citation_count", "引用数"),
    "monitor_rank": ("monitorRank", "monitor_rank", "监控位次"),
    "screenshot_url": ("screenshotUrl", "screenshot_url", "截图URL"),
    "sampled_at": ("captured_at", "采集时间", "监控时间", "date", "collectedAt"),
}
CITATION_ALIASES = {
    "citation_id": ("citation_id", "信源ID", "引用ID", "sourceRecordId", "source_record_id"),
    "question_id": ("questionId", "question_id", "问题ID"),
    "answer_id": ("answer_id", "答案ID", "回答ID", "observation_id", "sampleSourceRecordId"),
    "sample_source_record_id": ("sampleSourceRecordId", "sample_source_record_id", "样本信源记录ID"),
    "platform": ("platform", "平台", "engine"),
    "model": ("model", "模型", "model_name"),
    "title": ("title", "标题", "页面标题", "source_title"),
    "url": ("url", "URL", "链接", "引用链接", "source_url"),
    "publisher": ("media", "媒体", "来源名称", "publisher"),
    "domain": ("domain", "域名"),
    "published_at": ("publishedAt", "published_at", "发布时间"),
    "collected_at": ("collectedAt", "collected_at", "采集时间"),
    "author": ("author", "authorName", "author_name", "作者", "署名"),
    "accessed_at": ("accessedAt", "accessed_at", "访问时间", "页面访问时间", "检索时间"),
    "body_excerpt": ("body_excerpt", "bodyExcerpt", "source_excerpt", "page_excerpt", "正文摘录", "页面摘录"),
    "citation_context": (
        "citationContext", "citation_context", "citationPosition", "citation_position",
        "citation_text", "引用上下文", "引用位置", "引用片段", "引用内容", "snippet",
    ),
    "cited_claim": ("citedClaim", "cited_claim", "supportedClaim", "被引主张", "支撑主张", "引用观点"),
    "source_type": ("sourceType", "source_type", "信源类型", "来源类型"),
    "primary_source_status": (
        "primarySourceStatus", "primary_source_status", "isPrimarySource", "原始信源状态", "是否原始来源",
    ),
    "freshness_risk": ("freshnessRisk", "freshness_risk", "新鲜度风险", "时效风险"),
}
TRACKING_QUERY_KEYS = {
    "utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content",
    "gclid", "fbclid", "spm", "from", "source",
}


def _string(value: Any) -> str:
    return "" if value is None else str(value).strip()


def iso_datetime(value: str | None) -> str:
    raw = _string(value)
    if not raw:
        return datetime.now(timezone.utc).isoformat()
    normalized = raw.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(normalized)
    except ValueError:
        for fmt in ("%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(raw, fmt)
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"无法解析日期时间: {raw}")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.isoformat()


def optional_int(value: str, minimum: int = 0) -> int | None:
    if not _string(value):
        return None
    try:
        result = int(float(_string(value)))
    except ValueError:
        return None
    if result < minimum:
        raise ValueError(f"整数值 {result} 小于允许下限 {minimum}")
    return result


def normalize_enum(value: str, mapping: dict[str, str], allowed: set[str], field: str) -> str | None:
    raw = _string(value)
    if not raw:
        return None
    normalized = mapping.get(raw.strip().lower(), raw.strip().lower())
    if normalized not in allowed:
        raise ValueError(f"{field} 无效: {raw}")
    return normalized


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise RuntimeError("读取 XLSX 需要 openpyxl；不得把依赖缺失当作空表") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    try:
        headers = [_string(value) for value in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, row)) for row in rows if any(_string(value) for value in row)]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_input(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            payload = payload.get("rows", payload.get("data", []))
        if not isinstance(payload, list) or not all(isinstance(row, dict) for row in payload):
            raise ValueError(f"{path}: JSON 必须是对象数组或包含 rows/data 对象数组")
        if not payload:
            raise ValueError(f"{path}: 监控表必须至少包含一条有效记录")
        return payload
    if suffix in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(path)
    elif suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle, delimiter=delimiter))
    else:
        raise ValueError(f"不支持的监控表格式: {path.suffix}")
    if not rows:
        raise ValueError(f"{path}: 监控表必须至少包含一条有效记录")
    return rows


def read_rows(path: Path) -> list[dict[str, Any]]:
    return read_input(path)


def pick(row: dict[str, Any], aliases: Iterable[str]) -> str:
    folded = {str(key).strip().lower(): value for key, value in row.items()}
    for alias in aliases:
        value = folded.get(alias.lower())
        if _string(value):
            return _string(value)
    return ""


def canonical_url(raw: str) -> str:
    if not raw:
        return ""
    parts = urlsplit(raw)
    query = [(key, value) for key, value in parse_qsl(parts.query, keep_blank_values=True)
             if key.lower() not in TRACKING_QUERY_KEYS and not key.lower().startswith("utm_")]
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path.rstrip("/"), urlencode(query), ""))


def mapped(row: dict[str, Any], aliases: dict[str, tuple[str, ...]]) -> dict[str, str]:
    return {key: pick(row, values) for key, values in aliases.items()}


def validate_derived(
    derived: dict[str, Any], answer_ids: set[str], observation_ids: set[str],
    question_mismatch_answer_ids: set[str],
) -> None:
    required = {
        "competitors", "excluded_observations", "answer_landscape",
        "supplementary_primary_sources", "recurring_subquestions", "warnings",
    }
    missing, extras = sorted(required - set(derived)), sorted(set(derived) - required)
    if missing or extras:
        raise ValueError(f"derived 字段不匹配根 Schema；missing={missing}, extras={extras}")
    for key in required:
        if not isinstance(derived[key], list):
            raise ValueError(f"derived.{key} 必须是数组")
    known_records = answer_ids | observation_ids
    excluded_ids: set[str] = set()
    for item in derived["excluded_observations"]:
        if not isinstance(item, dict) or set(item) != {"record_id", "reason"} or not str(item.get("reason", "")).strip():
            raise ValueError("excluded_observations 每项必须包含 record_id 与 reason")
        record_id = str(item["record_id"])
        if record_id not in known_records:
            raise ValueError(f"excluded_observations 引用了未知记录: {record_id}")
        excluded_ids.add(record_id)
    if not (answer_ids - excluded_ids):
        raise ValueError("所有答案观察都与正式问题不同；必须阻断而不是继续制作")
    not_excluded = question_mismatch_answer_ids - excluded_ids
    if not_excluded:
        raise ValueError(f"题面与正式问题不一致的答案必须进入 excluded_observations: {sorted(not_excluded)}")

    landscapes = derived["answer_landscape"]
    if not landscapes:
        raise ValueError("derived.answer_landscape 必须至少覆盖一条有效答案")
    required_landscape = {"platform", "answer_id", "direct_answer", "entities", "stance", "dimensions", "unanswered_subquestions"}
    for item in landscapes:
        if not isinstance(item, dict) or set(item) != required_landscape:
            raise ValueError("answer_landscape 每项必须精确符合根 Schema 字段")
        if item["answer_id"] not in answer_ids:
            raise ValueError(f"answer_landscape 引用了未知 answer_id: {item['answer_id']}")
        if item["answer_id"] in excluded_ids:
            raise ValueError(f"answer_landscape 不得引用已排除答案: {item['answer_id']}")
        if item["stance"] not in {"positive", "neutral", "negative", "mixed", "not_applicable"}:
            raise ValueError(f"answer_landscape stance 无效: {item['stance']}")
        for key in ("entities", "dimensions", "unanswered_subquestions"):
            values = item[key]
            if (not isinstance(values, list) or len(values) != len(set(values))
                    or any(not isinstance(value, str) or not value.strip() for value in values)):
                raise ValueError(f"answer_landscape.{key} 必须是唯一非空字符串数组")
        if not str(item["platform"]).strip() or not str(item["direct_answer"]).strip():
            raise ValueError("answer_landscape platform/direct_answer 不能为空")
    for item in derived["competitors"]:
        if (not isinstance(item, dict) or set(item) != {"name", "answer_ids"}
                or not str(item.get("name", "")).strip() or not isinstance(item.get("answer_ids"), list)
                or not item["answer_ids"] or len(item["answer_ids"]) != len(set(item["answer_ids"]))):
            raise ValueError("competitors 每项必须包含 name 与非空 answer_ids")
        competitor_answer_ids = set(str(value) for value in item["answer_ids"])
        unknown = competitor_answer_ids - answer_ids
        if unknown:
            raise ValueError(f"competitor 引用了未知 answer_ids: {sorted(unknown)}")
        excluded = competitor_answer_ids & excluded_ids
        if excluded:
            raise ValueError(f"competitor 不得引用已排除答案: {sorted(excluded)}")
    competitor_names = [str(item.get("name", "")).strip() for item in derived["competitors"] if isinstance(item, dict)]
    if len(competitor_names) != len(set(competitor_names)):
        raise ValueError("competitors.name 必须唯一")
    required_supplementary = {"url", "title", "supports_subquestions", "selection_reason", "accessed_at"}
    allowed_supplementary = required_supplementary | {"publisher", "source_type"}
    for item in derived["supplementary_primary_sources"]:
        if (not isinstance(item, dict) or not required_supplementary.issubset(item)
                or set(item) - allowed_supplementary):
            raise ValueError("supplementary_primary_sources 缺少根 Schema 必需字段")
        if item.get("source_type") not in {None, "regulation_standard", "primary_data_research", "first_party_official", "authorized_project_record"}:
            raise ValueError("supplementary_primary_sources.source_type 无效")
        if not valid_nonempty_string_list(item.get("supports_subquestions"), require_one=True, unique=True):
            raise ValueError("supplementary_primary_sources.supports_subquestions 必须为非空唯一字符串数组")


def valid_nonempty_string_list(value: Any, require_one: bool = False, unique: bool = False) -> bool:
    if not isinstance(value, list) or (require_one and not value):
        return False
    if any(not isinstance(item, str) or not item.strip() for item in value):
        return False
    return not unique or len(value) == len(set(value))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--job-id", "--task-id", dest="job_id", required=True)
    parser.add_argument("--question", required=True)
    parser.add_argument("--answers", type=Path, required=True)
    parser.add_argument("--citations", type=Path, required=True)
    parser.add_argument("--derived-analysis", type=Path, required=True,
                        help="只含答案格局/竞品/补充一手来源的 derived JSON；不得含Top20/P分类")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--raw-metadata-output", type=Path)
    parser.add_argument("--package-root", type=Path, required=True,
                        help="staged job package root; input/output paths must live inside it")
    args = parser.parse_args()
    for path in (args.answers, args.citations, args.derived_analysis):
        if not path.is_file():
            raise FileNotFoundError(path)
    package_root = args.package_root.resolve()
    if not package_root.is_dir():
        raise FileNotFoundError(package_root)

    def package_path(path: Path) -> str:
        try:
            return path.resolve().relative_to(package_root).as_posix()
        except ValueError as exc:
            raise ValueError(f"路径必须先暂存到 --package-root 内，避免泄露本机绝对路径: {path}") from exc

    # Validate staging before doing any analysis or writing partial output.
    package_path(args.answers)
    package_path(args.citations)
    package_path(args.derived_analysis)
    package_path(args.output)

    answer_rows = read_input(args.answers)
    citation_rows = read_input(args.citations)
    raw_answers = [mapped(row, ANSWER_ALIASES) for row in answer_rows]
    raw_citations = [mapped(row, CITATION_ALIASES) for row in citation_rows]

    citations_by_answer: dict[str, list[dict[str, Any]]] = {}
    source_observations = []
    seen_urls: set[str] = set()
    seen_observation_ids: set[str] = set()
    warnings: list[str] = []
    for index, item in enumerate(raw_citations, 1):
        if not item["url"]:
            raise ValueError(f"信源表第 {index + 1} 行缺少 url")
        citation_id = item["citation_id"] or f"cit_{index:04d}"
        url = canonical_url(item["url"])
        parsed_url = urlsplit(url)
        if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
            raise ValueError(f"信源表第 {index + 1} 行 URL 必须是完整 HTTP(S) 地址: {item['url']}")
        answer_id = item["answer_id"] or item["sample_source_record_id"]
        citation = {
            "citation_id": citation_id,
            "url": url,
            "title": item["title"] or None,
            "publisher": item["publisher"] or item["domain"] or None,
        }
        if answer_id:
            citations_by_answer.setdefault(answer_id, []).append(citation)
        else:
            warnings.append(f"信源 {citation_id} 未提供 sampleSourceRecordId/answer_id，保留为未关联 source observation")
        if url in seen_urls:
            continue
        if citation_id in seen_observation_ids:
            raise ValueError(f"信源表存在重复 observation_id 且 URL 不同: {citation_id}")
        seen_urls.add(url)
        seen_observation_ids.add(citation_id)
        source_type = normalize_enum(
            item["source_type"],
            {
                "法规/标准": "regulation_standard", "法规标准": "regulation_standard",
                "原始数据/研究": "primary_data_research", "原始研究": "primary_data_research",
                "官网": "first_party_official", "一手官方": "first_party_official",
                "授权项目记录": "authorized_project_record", "媒体": "editorial_media",
                "专业数据库": "professional_database", "专家评论": "expert_commentary",
                "论坛/评价": "forum_review", "聚合页": "aggregator", "未知": "unknown",
            },
            {
                "regulation_standard", "primary_data_research", "first_party_official", "authorized_project_record",
                "editorial_media", "professional_database", "expert_commentary", "forum_review", "aggregator", "unknown",
            }, "source_type",
        )
        primary_status = normalize_enum(
            item["primary_source_status"],
            {
                "true": "primary", "yes": "primary", "1": "primary", "是": "primary", "原始": "primary", "一手": "primary",
                "false": "secondary", "no": "secondary", "0": "secondary", "否": "secondary", "二手": "secondary",
                "不确定": "uncertain", "待确认": "uncertain",
            }, {"primary", "secondary", "uncertain"}, "primary_source_status",
        )
        freshness_risk = normalize_enum(
            item["freshness_risk"],
            {"无": "none", "低": "low", "中": "medium", "高": "high", "未知": "unknown"},
            {"none", "low", "medium", "high", "unknown"}, "freshness_risk",
        )
        source_observations.append({
            "observation_id": citation_id,
            "source_record_id": item["citation_id"] or None,
            "question_id": item["question_id"] or None,
            "sample_source_record_id": item["sample_source_record_id"] or None,
            "model": item["model"] or None,
            "media": item["publisher"] or None,
            "domain": item["domain"] or None,
            "collected_at": iso_datetime(item["collected_at"]) if item["collected_at"] else None,
            "url": url,
            "retrieval_status": "not_attempted",
            "title": item["title"] or None,
            "publisher": item["publisher"] or item["domain"] or None,
            "author": item["author"] or None,
            "published_at": iso_datetime(item["published_at"]) if item["published_at"] else None,
            "accessed_at": iso_datetime(item["accessed_at"]) if item["accessed_at"] else None,
            "body_excerpt": item["body_excerpt"] or None,
            "citation_context": item["citation_context"] or None,
            "cited_claim": item["cited_claim"] or None,
            "source_type": source_type,
            "primary_source_status": primary_status,
            "freshness_risk": freshness_risk,
            "question_similarity": None,
        })

    answers = []
    known_answer_ids: set[str] = set()
    question_mismatch_answer_ids: set[str] = set()
    for index, item in enumerate(raw_answers, 1):
        if not item["answer_text"] or not item["platform"]:
            raise ValueError(f"答案表第 {index + 1} 行缺少 platform/content")
        question = item["question"] or args.question.strip()
        answer_id = item["answer_id"] or f"ans_{index:04d}"
        if answer_id in known_answer_ids:
            raise ValueError(f"答案表存在重复 answer_id: {answer_id}")
        known_answer_ids.add(answer_id)
        if question != args.question.strip():
            question_mismatch_answer_ids.add(answer_id)
            warnings.append(f"答案 {answer_id} 题面与正式问题不完全一致，已要求进入 excluded_observations")
        platform = item["platform"]
        if item["model"] and item["model"].lower() not in platform.lower():
            platform = f"{platform} / {item['model']}"
        answers.append({
            "answer_id": answer_id,
            "platform": platform,
            "question_id": item["question_id"] or None,
            "model": item["model"] or None,
            "answer_no": optional_int(item["answer_no"], minimum=1),
            "citation_count": optional_int(item["citation_count"], minimum=0),
            "monitor_rank": optional_int(item["monitor_rank"], minimum=1),
            "screenshot_url": item["screenshot_url"] or None,
            "source_record_id": item["answer_id"] or None,
            "collected_at": iso_datetime(item["sampled_at"]) if item["sampled_at"] else None,
            "sampled_at": iso_datetime(item["sampled_at"]),
            "answer_text": item["answer_text"],
            "citations": citations_by_answer.get(answer_id, []),
        })
    dangling = sorted(set(citations_by_answer) - known_answer_ids)
    if dangling:
        warnings.append(f"存在未关联到答案表 answer_id 的信源引用: {', '.join(dangling)}")
    if not answers or not source_observations:
        raise ValueError("答案表与信源表规范化后都必须至少保留一条有效记录")

    derived = json.loads(args.derived_analysis.read_text(encoding="utf-8"))
    if not isinstance(derived, dict):
        raise ValueError("--derived-analysis 必须包含 JSON object")
    observation_ids = {str(item["observation_id"]) for item in source_observations}
    validate_derived(derived, known_answer_ids, observation_ids, question_mismatch_answer_ids)
    derived["warnings"] = list(dict.fromkeys([*derived["warnings"], *warnings]))
    status = "partial" if derived["warnings"] or derived["excluded_observations"] else "available"
    captured_candidates = [item["sampled_at"] for item in answers]
    captured_candidates += [iso_datetime(item["collected_at"]) for item in raw_citations if item["collected_at"]]
    captured_at = max(captured_candidates) if captured_candidates else datetime.now(timezone.utc).isoformat()
    metadata_path = args.raw_metadata_output or args.output.with_name(args.output.stem + ".raw_metadata.json")
    payload = {
        "schema_version": "2.1.0",
        "job_id": args.job_id,
        "question": args.question.strip(),
        "captured_at": captured_at,
        "input_files": {
            "answers_table_path": package_path(args.answers),
            "sources_table_path": package_path(args.citations),
            "answers_table_sha256": file_sha256(args.answers),
            "sources_table_sha256": file_sha256(args.citations),
            "normalization_sidecar_path": package_path(metadata_path),
        },
        "answers": answers,
        "source_observations": source_observations,
        "analysis_status": status,
        "derived": derived,
    }
    validate_root(payload, "monitoring_input.schema.json", "E2 monitoring_input")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    metadata_path.parent.mkdir(parents=True, exist_ok=True)
    metadata_path.write_text(json.dumps({
        "schema_version": "2.1.0",
        "job_id": args.job_id,
        "raw_answer_rows": answer_rows,
        "raw_source_rows": citation_rows,
        "normalized_answer_columns": raw_answers,
        "normalized_source_columns": raw_citations,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": package_path(args.output), "raw_metadata": package_path(metadata_path), "analysis_status": status,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
