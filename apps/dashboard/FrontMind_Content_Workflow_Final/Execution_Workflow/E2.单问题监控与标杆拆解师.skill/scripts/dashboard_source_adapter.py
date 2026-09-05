#!/usr/bin/env python3
"""Tolerant Dashboard source-workbook adapter for FrontMind E2 v2.3.

Mixed-question exports are filtered to the current task.  Header aliases,
broken ranks and stale citation-share cells are repaired deterministically and
reported as warnings instead of stopping content production.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import tempfile
import unicodedata
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook


SHEET_ALIASES = {
    "statistics": ("内容统计", "内容汇总", "信源统计", "contentstatistics", "contentstats"),
    "details": ("详细表格", "引用明细", "信源明细", "details", "detailedtable"),
}
HEADER_ALIASES = {
    "rank": ("引用排行", "排行", "排名", "rank", "citationrank"),
    "title": ("内容标题", "标题", "文章标题", "title", "contenttitle"),
    "url": ("内容url", "内容链接", "链接", "文章链接", "url", "contenturl"),
    "media": ("媒体名称", "媒体", "来源", "publisher", "medianame"),
    "media_url": ("媒体url", "媒体链接", "域名", "mediaurl", "domain"),
    "date": ("引用日期", "日期", "时间", "citationdate", "date"),
    "count": ("引用次数", "被引次数", "次数", "citationcount", "count"),
    "share": ("引用占比", "占比", "citationshare", "share"),
    "question": ("监控问题", "问题", "正式问题", "question", "query"),
}
TRACKING_KEYS = {"utm_source", "utm_medium", "utm_campaign", "utm_term", "utm_content", "gclid", "fbclid", "spm"}


def normal(value: Any) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value or ""))).strip()


def key(value: Any) -> str:
    return re.sub(r"[^0-9a-z\u3400-\u9fff]+", "", unicodedata.normalize("NFKC", normal(value)).casefold())


def question_key(value: Any) -> str:
    """Normalize spelling and punctuation without guessing semantic similarity."""
    return re.sub(r"[^0-9a-z\u3400-\u9fff]+", "", unicodedata.normalize("NFKC", normal(value)).casefold())


def canonical_url(value: Any) -> str:
    raw = normal(value)
    if not raw:
        return ""
    try:
        parsed = urlsplit(raw)
        if parsed.scheme.casefold() not in {"http", "https"} or not parsed.netloc:
            return raw
        hostname = (parsed.hostname or "").casefold()
        port = parsed.port
        netloc = hostname
        if port and not ((parsed.scheme == "http" and port == 80) or (parsed.scheme == "https" and port == 443)):
            netloc = f"{hostname}:{port}"
        query = urlencode(sorted(
            (name, item) for name, item in parse_qsl(parsed.query, keep_blank_values=True)
            if name.casefold() not in TRACKING_KEYS
        ), doseq=True)
        return urlunsplit((parsed.scheme.casefold(), netloc, parsed.path or "/", query, ""))
    except ValueError:
        return raw


def number(value: Any, default: int = 0) -> int:
    if isinstance(value, bool):
        return default
    try:
        return max(0, int(float(value)))
    except (TypeError, ValueError):
        return default


def share_value(value: Any) -> float | None:
    raw = normal(value)
    if not raw:
        return None
    try:
        result = float(raw[:-1]) / 100 if raw.endswith("%") else float(raw)
        if result > 1:
            result /= 100
        return result if 0 <= result <= 1 else None
    except (TypeError, ValueError):
        return None


def date_text(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normal(value)
    return text or None


def domain(value: Any, fallback_url: str) -> str | None:
    raw = normal(value)
    candidate = raw if "://" in raw else f"https://{raw}" if raw else fallback_url
    try:
        return (urlsplit(candidate).hostname or "").casefold() or None
    except ValueError:
        return None


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def choose_sheet(workbook: Any, kind: str) -> Any | None:
    aliases = set(SHEET_ALIASES[kind])
    for name in workbook.sheetnames:
        if key(name) in aliases:
            return workbook[name]
    return None


def headers(sheet: Any) -> dict[str, int]:
    first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    normalized = [key(value) for value in first]
    result: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            normalized_alias = key(alias)
            if normalized_alias in normalized:
                result[field] = normalized.index(normalized_alias)
                break
    return result


def cell(values: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns.get(name)
    return values[index] if index is not None and index < len(values) else None


def row_objects(sheet: Any | None) -> tuple[list[dict[str, Any]], dict[str, int]]:
    if sheet is None:
        return [], {}
    columns = headers(sheet)
    result: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values):
            continue
        result.append({
            "row": row_number,
            **{name: cell(values, columns, name) for name in HEADER_ALIASES},
        })
    return result, columns


def choose_question(rows: list[dict[str, Any]], target: str, warnings: list[str]) -> tuple[list[dict[str, Any]], str, str]:
    questions = Counter(normal(row.get("question")) for row in rows if normal(row.get("question")))
    if not questions:
        raise ValueError("research_question_unidentifiable: 详细表格没有可识别的监控问题列")
    normalized_target = normal(target)
    target_key = question_key(normalized_target)
    exact_variants = {question for question in questions if question_key(question) == target_key}
    if exact_variants:
        if len(questions) > 1:
            warnings.append(f"工作簿包含 {len(questions)} 个问题，已自动筛选正式问题")
        return [row for row in rows if normal(row.get("question")) in exact_variants], normalized_target, "exact_normalized"
    choices = " | ".join(question for question, _count in questions.most_common(10))
    raise ValueError(
        f"research_question_mismatch: 信源工作簿没有精确对应正式问题 {normalized_target!r} 的行；"
        f"检测到：{choices}"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a repairable E2 Top20 pool")
    parser.add_argument("--content-job", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    job = json.loads(args.content_job.read_text(encoding="utf-8"))
    if not isinstance(job, dict):
        raise ValueError("content job must be an object")
    task = job.get("task") if isinstance(job.get("task"), dict) else {}
    question = normal(task.get("question_text"))
    if not question or job.get("entry_category") == "foundation_start":
        raise ValueError("foundation_start does not run E2; ordinary entries require a question")
    if args.workbook.is_symlink() or not args.workbook.is_file() or args.workbook.suffix.casefold() != ".xlsx":
        raise ValueError("source workbook must be a regular .xlsx file")

    warnings: list[str] = []
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        statistics_sheet = choose_sheet(workbook, "statistics")
        detail_sheet = choose_sheet(workbook, "details")
        if statistics_sheet is None:
            warnings.append("缺少内容统计页，已从详细表格重建排行")
        if detail_sheet is None:
            raise ValueError("research_question_unidentifiable: 缺少可按正式问题筛选的详细表格页")
        statistics, statistics_columns = row_objects(statistics_sheet)
        details, detail_columns = row_objects(detail_sheet)
        selected_details, selected_question, selection_method = choose_question(details, question, warnings)
        if statistics_sheet is not None and "url" not in statistics_columns:
            warnings.append("内容统计没有标准 URL 列，无法采用该页")
            statistics = []
        if detail_sheet is not None and "url" not in detail_columns:
            raise ValueError("source_workbook_unusable: 详细表格没有可识别的 URL 列")

        detail_by_url: dict[str, list[dict[str, Any]]] = {}
        for row in selected_details:
            url = canonical_url(row.get("url"))
            if url:
                detail_by_url.setdefault(url, []).append(row)
        stats_by_url: dict[str, dict[str, Any]] = {}
        for order, row in enumerate(statistics, 1):
            url = canonical_url(row.get("url"))
            if not url:
                continue
            stored = dict(row)
            stored["original_order"] = order
            stats_by_url.setdefault(url, stored)

        use_detail_counts = bool(detail_by_url)
        urls = list(detail_by_url) if use_detail_counts else list(stats_by_url)
        candidates: list[dict[str, Any]] = []
        for order, url in enumerate(urls, 1):
            detail_rows = detail_by_url.get(url, [])
            stats = stats_by_url.get(url, {})
            exemplar = detail_rows[0] if detail_rows else stats
            count = len(detail_rows) if detail_rows else number(stats.get("count"), 0)
            if count <= 0:
                count = 1
                warnings.append(f"引用次数不可用，已按一次引用处理：{url}")
            exported_rank = number(stats.get("rank"), 0) or None
            candidates.append({
                "original_order": number(stats.get("original_order"), order),
                "exported_rank": exported_rank,
                "content_title": normal(exemplar.get("title")) or normal(stats.get("title")) or url,
                "canonical_url": url,
                "media_name": normal(exemplar.get("media")) or normal(stats.get("media")) or None,
                "media_domain": domain(exemplar.get("media_url") or stats.get("media_url"), url),
                "citation_date": date_text(exemplar.get("date") or stats.get("date")),
                "citation_count": count,
                "exported_share": share_value(stats.get("share")),
                "detailed_row_refs": [int(row["row"]) for row in detail_rows],
            })

        exported_ranks = [item["exported_rank"] for item in candidates if item["exported_rank"]]
        rank_is_clean = (
            len(exported_ranks) == len(candidates)
            and len(exported_ranks) == len(set(exported_ranks))
            and sorted(exported_ranks) == list(range(1, len(candidates) + 1))
        )
        counts_by_export_rank = [
            item["citation_count"] for item in sorted(candidates, key=lambda item: item["exported_rank"] or 10**9)
        ]
        descending = all(left >= right for left, right in zip(counts_by_export_rank, counts_by_export_rank[1:]))
        if use_detail_counts or not rank_is_clean or not descending:
            candidates.sort(key=lambda item: (-item["citation_count"], item["exported_rank"] or 10**9, item["original_order"], item["canonical_url"]))
            if candidates and (not rank_is_clean or not descending):
                warnings.append("引用排行存在重复、断号或与引用次数不一致，已按引用次数稳定重排")
            elif use_detail_counts and len(Counter(normal(row.get("question")) for row in details if normal(row.get("question")))) > 1:
                warnings.append("混合问题导出已按筛选后的详细表引用次数重建排行")
        else:
            candidates.sort(key=lambda item: item["exported_rank"])

        # Citation share keeps the Dashboard definition: this source's count
        # divided by every citation for the selected question, not by the
        # Top20 subtotal.  Only the displayed pool is truncated to twenty.
        total = sum(item["citation_count"] for item in candidates) or 1
        candidates = candidates[:20]
        pool: list[dict[str, Any]] = []
        repaired_share = False
        for rank, item in enumerate(candidates, 1):
            calculated = item["citation_count"] / total
            if item["exported_share"] is None or abs(item["exported_share"] - calculated) > 0.000051:
                repaired_share = True
            pool.append({
                "raw_rank": rank,
                "content_title": item["content_title"],
                "canonical_url": item["canonical_url"],
                "media_name": item["media_name"],
                "media_domain": item["media_domain"],
                "citation_date": item["citation_date"],
                "citation_count": item["citation_count"],
                "citation_share": round(calculated, 12),
                "detailed_row_refs": item["detailed_row_refs"],
            })
        if repaired_share and pool:
            warnings.append("引用占比缺失或与筛选后的引用量不一致，已自动重算")
        if not pool:
            raise ValueError("source_workbook_unusable: 精确问题范围内没有可识别的引用 URL")

        payload = {
            "schema_version": "2.3.0",
            "artifact_type": "E2_source_pool",
            "job_id": job.get("job_id"),
            "question": question,
            "entry_category": job.get("entry_category"),
            "product_scenario_subintent": task.get("product_scenario_subintent"),
            "created_at": datetime.now(timezone.utc).isoformat(),
            "workbook_summary": {
                "sheet_names": list(workbook.sheetnames),
                "selected_question": selected_question,
                "question_selection_method": selection_method,
                "selected_detail_rows": len(selected_details),
            },
            "pool_status": "complete_pool" if len(pool) == 20 else "partial_pool" if pool else "not_available",
            "cited_content_pool": pool,
            "warnings": list(dict.fromkeys(warnings)),
        }
    finally:
        workbook.close()

    write_json(args.output, payload)
    print(json.dumps({
        "stage": "E2", "status": "completed_with_limits" if warnings else "completed",
        "pool_status": payload["pool_status"], "pool_size": len(payload["cited_content_pool"]),
        "warnings": payload["warnings"], "output": str(args.output),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
