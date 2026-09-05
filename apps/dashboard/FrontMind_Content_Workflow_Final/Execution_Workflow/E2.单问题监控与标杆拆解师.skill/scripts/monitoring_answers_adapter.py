#!/usr/bin/env python3
"""Normalize Dashboard AI-answer exports into the E2 v2.3 monitoring context.

The adapter consumes XLSX, CSV and JSON directly.  It deliberately keeps
answer observations separate from citation-page research: answer mentions are
candidate and question signals, not proof of ranking or competitor claims.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared"))
from schema_validation import validate_root  # noqa: E402


QUESTION_ALIASES = {"question", "query", "prompt", "monitoringquestion", "monitorquestion", "问题", "监控问题", "正式问题"}
PLATFORM_ALIASES = {"platform", "provider", "channel", "平台", "模型平台"}
MODEL_ALIASES = {"model", "modelname", "模型", "模型名称"}
DATE_ALIASES = {"date", "datetime", "time", "capturedat", "sampledat", "日期", "时间", "监控时间"}
ANSWER_ALIASES = {"answer", "answertext", "content", "response", "output", "答案", "答案内容", "内容", "回答"}
URL_RE = re.compile(r"https?://[^\s<>\]\[\"'、，。；！？)]+", re.I)
QUESTION_TRAIL_RE = re.compile(r"([^\n。！!?]{3,120}[?？])")
ORG_SUFFIXES = (
    "医疗美容医院|整形美容医院|医疗美容门诊部|整形美容门诊部|"
    "医学美容中心|医疗美容中心|美容医院|医院|门诊部|诊所|中心|"
    "医疗美容|整形美容|医美|整形|机构|品牌|公司|科技|大学|学院"
)
ORG_RE = re.compile(rf"[一-鿿A-Za-z0-9·・()（）\-]{{2,36}}?(?:{ORG_SUFFIXES})")
BOLD_RE = re.compile(r"\*\*([^*\n]{2,40})\*\*")
ENTITY_STOP = {
    "医疗机构", "正规机构", "其他机构", "客户品牌", "相关品牌", "官方网站",
    "注射微整", "玻尿酸", "选择建议", "注意事项", "就诊小贴士", "风险提示", "核心结论",
}
NON_ENTITY_CUES = {
    "擅长", "项目", "技术", "特点", "优势", "服务流程", "注意", "推荐理由",
    "选择建议", "参考资料", "官方信息", "口碑", "安全", "风险", "资质核验",
    "产品验真", "费用构成", "异常处置", "适用人群", "核心结论", "温馨提示",
}
DIMENSION_CUES = {
    "机构与医生资质": ("资质", "执业许可", "执业证", "主诊", "医师"),
    "产品与设备来源": ("产品", "正品", "批次", "NMPA", "注册证", "设备", "验真"),
    "专业能力与技术": ("技术", "擅长", "专业", "经验", "方案", "审美"),
    "需求与适用情境": ("适合", "需求", "场景", "适用", "面诊", "人群"),
    "费用与价格构成": ("价格", "费用", "报价", "收费", "预算"),
    "流程与交付": ("流程", "交付", "预约", "步骤", "使用", "服务"),
    "风险与安全": ("风险", "安全", "并发症", "急救", "异常", "禁忌"),
    "复诊与售后": ("复诊", "随访", "售后", "术后", "保修", "支持"),
    "口碑与公开反馈": ("口碑", "评价", "投诉", "反馈", "信誉"),
}


def normal(value: Any) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value or ""))).strip()


def key(value: Any) -> str:
    return re.sub(r"[^0-9a-z㐀-鿿]+", "", unicodedata.normalize("NFKC", normal(value)).casefold())


def question_key(value: Any) -> str:
    return re.sub(r"[^0-9a-z㐀-鿿]+", "", unicodedata.normalize("NFKC", normal(value)).casefold())


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def iso_datetime(value: Any) -> str | None:
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.isoformat()
    if isinstance(value, date):
        return f"{value.isoformat()}T00:00:00+08:00"
    raw = normal(value)
    if not raw:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
        return f"{raw}T00:00:00+08:00"
    candidate = raw.replace(" ", "T", 1)
    try:
        parsed = datetime.fromisoformat(candidate)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.isoformat()


def valid_public_url(value: Any) -> str | None:
    raw = normal(value).rstrip(".,;，。；")
    try:
        parsed = urlsplit(raw)
    except ValueError:
        return None
    return raw if parsed.scheme.casefold() in {"http", "https"} and parsed.netloc else None


def citations_from_text(text: str) -> list[dict[str, Any]]:
    result = []
    for raw in URL_RE.findall(text):
        url = valid_public_url(raw)
        if url and url not in {item["url"] for item in result}:
            result.append({"url": url, "title": None, "publisher": urlsplit(url).hostname})
    return result


def field(row: dict[str, Any], aliases: set[str]) -> Any:
    for name, value in row.items():
        if key(name) in aliases and value not in (None, ""):
            return value
    return None


def answer_fields(row: dict[str, Any]) -> list[tuple[str, Any]]:
    result: list[tuple[str, Any]] = []
    for name, value in row.items():
        name_key = key(name)
        if value in (None, ""):
            continue
        if name_key in ANSWER_ALIASES or (
            (name_key.startswith("answer") or name_key.startswith("答案"))
            and (name_key.endswith("content") or name_key.endswith("内容") or re.fullmatch(r"(?:answer|答案)\d*", name_key))
        ):
            result.append((normal(name), value))
    return result


def rows_from_csv(path: Path) -> list[tuple[str, dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        return [(path.stem, dict(row)) for row in csv.DictReader(handle)]


def composed_headers(first: tuple[Any, ...], second: tuple[Any, ...] | None) -> tuple[list[str], int]:
    first_names = [normal(item) for item in first]
    if second is None:
        return [name or f"column_{index + 1}" for index, name in enumerate(first_names)], 2
    second_names = [normal(item) for item in second]
    two_rows = any(name in {"内容", "content", "截图链接", "排名", "监控词排名"} for name in second_names)
    if not two_rows:
        return [name or f"column_{index + 1}" for index, name in enumerate(first_names)], 2
    headers: list[str] = []
    for index in range(max(len(first_names), len(second_names))):
        top = first_names[index] if index < len(first_names) else ""
        bottom = second_names[index] if index < len(second_names) else ""
        headers.append(f"{top}_{bottom}" if top and bottom and key(top) != key(bottom) else top or bottom or f"column_{index + 1}")
    return headers, 3


def rows_from_xlsx(path: Path) -> list[tuple[str, dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - runtime dependency error
        raise RuntimeError("openpyxl is required to read monitoring XLSX files") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    result: list[tuple[str, dict[str, Any]]] = []
    try:
        for sheet in workbook.worksheets:
            raw_rows = sheet.iter_rows(values_only=True)
            first = next(raw_rows, ())
            second = next(raw_rows, None)
            headers, start_row = composed_headers(first, second)
            values_iter: Iterable[tuple[Any, ...]]
            if start_row == 2 and second is not None:
                values_iter = [second, *raw_rows]
            else:
                values_iter = raw_rows
            for values in values_iter:
                if not any(value not in (None, "") for value in values):
                    continue
                result.append((normal(sheet.title) or path.stem, {
                    headers[index]: value for index, value in enumerate(values) if index < len(headers)
                }))
    finally:
        workbook.close()
    return result


def rows_from_json(path: Path) -> tuple[list[tuple[str, dict[str, Any]]], dict[str, Any] | None]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(value, dict) and isinstance(value.get("answers"), list):
        return [], value
    raw_rows = value if isinstance(value, list) else value.get("rows") if isinstance(value, dict) else None
    if not isinstance(raw_rows, list):
        raise ValueError("monitoring JSON must be a normalized context, an answer array, or an object with rows[]")
    return [(path.stem, item) for item in raw_rows if isinstance(item, dict)], None


def extract_entities(text: str, question: str = "") -> list[str]:
    ordered: list[str] = []
    institution_question = not any(token in question for token in ("医生", "医师", "专家")) and any(
        token in question for token in ("机构", "医院", "医美", "品牌", "公司", "企业")
    )
    matches = [(item, "organization") for item in ORG_RE.findall(text)]
    matches.extend((item, "bold") for item in BOLD_RE.findall(text))
    for match, origin in matches:
        candidate = normal(re.sub(r"^[\d一-鿿]{0,3}[.、:：\-]\s*", "", match)).strip("#* ：:。，,()（）")
        candidate = re.sub(
            r"^(?:(?:可以|可|还可)?(?:重点)?(?:关注|了解|考虑|选择|推荐|提及|比较|对比)|包括|例如|如|其中|以及|另外|还有|和|与|及)+",
            "", candidate,
        ).strip()
        if not (2 <= len(candidate) <= 40) or candidate in ENTITY_STOP:
            continue
        if any(token in candidate for token in NON_ENTITY_CUES | {"提示", "建议", "注意", "风险", "价格", "资质", "如何", "哪些", "什么"}):
            continue
        if institution_question and (
            candidate.startswith(("希望", "更看重", "想比较", "准备比较", "如果比较"))
            or any(token in candidate for token in ("页面公开", "资料显示", "公开资料", "可先了解", "放在同一轮"))
        ):
            continue
        if institution_question and candidate.endswith("医生"):
            continue
        if origin == "bold" and institution_question and not re.search(rf"(?:{ORG_SUFFIXES})$", candidate):
            continue
        if generic_entity(candidate, question):
            continue
        if candidate not in ordered:
            ordered.append(candidate)
    return ordered[:30]


def question_location(question: str) -> str:
    match = re.match(r"^([\u3400-\u9fff]{2,6}?)(?:市)?(?:打|做|注射|哪些|哪几家|哪里)", normal(question))
    return match.group(1) if match else ""


def generic_entity(name: str, question: str) -> bool:
    compact = key(name)
    location = key(question_location(question))
    if location and compact.startswith(location):
        compact = compact[len(location):]
    compact = re.sub(r"^(?:当地|本地|正规|专业|其他|推荐)+", "", compact)
    generic = {
        "机构", "医院", "诊所", "中心", "医美", "医美机构", "医疗机构",
        "医疗美容机构", "整形美容机构", "美容医院", "整形医院",
        "玻尿酸机构", "玻尿酸医院", "打玻尿酸机构", "注射玻尿酸机构",
    }
    return compact in generic


def entity_alias_key(name: str, question: str) -> str:
    compact = key(name)
    compact = compact.replace("医疗美容", "医美").replace("医学美容", "医美")
    compact = compact.replace("整形美容", "医美")
    compact = re.sub(r"(?:医院|门诊部|诊所|中心|机构|公司)$", "", compact)
    location = key(question_location(question))
    if location and compact.startswith(location) and len(compact) - len(location) >= 3:
        compact = compact[len(location):]
    return compact


def preferred_entity_name(names: list[str]) -> str:
    def score(name: str) -> tuple[int, int, int, str]:
        formal = bool(re.search(r"(?:医院|门诊部|诊所|中心|公司)$", name))
        expanded = any(token in name for token in ("医疗美容", "医学美容", "整形美容"))
        return int(formal), int(expanded), len(name), name
    return max(names, key=score)


def collapse_entity_aliases(names: Iterable[str], question: str) -> dict[str, str]:
    """Fold only strong spelling/abbreviation aliases; never fuzzy-merge brands."""
    clusters: dict[str, list[str]] = defaultdict(list)
    for name in names:
        alias = entity_alias_key(name, question)
        if len(alias) >= 3:
            clusters[alias].append(name)
        else:
            clusters[f"literal:{key(name)}"].append(name)
    return {
        name: preferred_entity_name(cluster)
        for cluster in clusters.values()
        for name in cluster
    }


def dimensions(text: str) -> list[str]:
    sample = normal(text).casefold()
    return [name for name, cues in DIMENSION_CUES.items() if any(cue.casefold() in sample for cue in cues)]


def stance(text: str) -> str:
    sample = normal(text)
    positive = sum(sample.count(item) for item in ("推荐", "值得", "优势", "可考虑", "适合"))
    negative = sum(sample.count(item) for item in ("不推荐", "争议", "处罚", "投诉", "风险", "不适合"))
    if positive and negative:
        return "mixed"
    if positive:
        return "positive"
    if negative:
        return "negative"
    return "neutral"


def direct_answer(text: str) -> str:
    clean = normal(re.sub(r"[#*_>`]+", " ", text))
    sentences = re.split(r"(?<=[。！？.!?])\s*", clean)
    return normal("".join(sentences[:2]))[:500] or clean[:500]


def questions_in(text: str) -> tuple[list[str], list[str]]:
    all_questions = [normal(item) for item in QUESTION_TRAIL_RE.findall(text)]
    all_questions = list(dict.fromkeys(item for item in all_questions if 4 <= len(item) <= 120))
    unanswered = [
        item for item in all_questions
        if any(cue in item for cue in ("要不要", "需要我", "是否需要", "想确认", "方便告诉", "想了解"))
    ]
    return all_questions, unanswered


def normalize_existing(value: dict[str, Any], job_id: str, target_question: str) -> dict[str, Any]:
    supplied_question = normal(value.get("question"))
    if supplied_question and question_key(supplied_question) != question_key(target_question):
        raise ValueError(f"research_question_mismatch: monitoring question is {supplied_question!r}, expected {target_question!r}")
    answers = [item for item in value.get("answers", []) if isinstance(item, dict) and normal(item.get("answer_text"))]
    if not answers:
        raise ValueError("monitoring_answers_empty: no readable answer text")
    # Rebuild even a nominally v2.3 object so legacy SHA/receipt fields and
    # out-of-contract answer properties never leak into the active context.
    return build_context(job_id, target_question, answers)


def build_context(job_id: str, question: str, raw_answers: list[dict[str, Any]]) -> dict[str, Any]:
    answers: list[dict[str, Any]] = []
    landscapes: list[dict[str, Any]] = []
    source_observations: list[dict[str, Any]] = []
    sources_seen: set[str] = set()
    entity_answers: dict[str, list[str]] = defaultdict(list)
    entity_positions: dict[str, list[int]] = defaultdict(list)
    dimension_counts: Counter[str] = Counter()
    subquestion_counts: Counter[str] = Counter()
    unanswered: list[str] = []
    for index, raw in enumerate(raw_answers, 1):
        text = normal(raw.get("answer_text") or raw.get("content") or raw.get("answer"))
        if not text:
            continue
        answer_id = normal(raw.get("answer_id")) or f"answer_{index:03d}"
        platform = normal(raw.get("platform")) or "unknown_platform"
        model = normal(raw.get("model")) or None
        sampled = iso_datetime(raw.get("sampled_at") or raw.get("date"))
        citations: list[dict[str, Any]] = []
        for citation in raw.get("citations", []) if isinstance(raw.get("citations"), list) else []:
            if not isinstance(citation, dict):
                continue
            url = valid_public_url(citation.get("url"))
            if url:
                citations.append({
                    "url": url,
                    "title": normal(citation.get("title")) or None,
                    "publisher": normal(citation.get("publisher")) or urlsplit(url).hostname,
                })
        for citation in citations_from_text(text):
            if citation["url"] not in {item["url"] for item in citations}:
                citations.append(citation)
        entities = extract_entities(text, question)
        answer_dimensions = dimensions(text)
        all_questions, answer_unanswered = questions_in(text)
        for position, entity in enumerate(entities, 1):
            entity_answers[entity].append(answer_id)
            entity_positions[entity].append(position)
        dimension_counts.update(answer_dimensions)
        subquestion_counts.update(all_questions)
        unanswered.extend(answer_unanswered)
        answers.append({
            "answer_id": answer_id,
            "platform": platform,
            "model": model,
            "sampled_at": sampled,
            "answer_text": text,
            "citations": citations,
        })
        landscapes.append({
            "answer_id": answer_id,
            "platform": platform,
            "direct_answer": direct_answer(text),
            "entities": entities,
            "stance": stance(text),
            "dimensions": answer_dimensions,
            "unanswered_subquestions": answer_unanswered,
        })
        for citation in citations:
            if citation["url"] not in sources_seen:
                sources_seen.add(citation["url"])
                source_observations.append(citation)
    if not answers:
        raise ValueError("monitoring_answers_empty: no readable answer text")
    aliases = collapse_entity_aliases(entity_answers, question)
    entity_answers = defaultdict(list)
    entity_positions = defaultdict(list)
    for landscape in landscapes:
        folded_entities = list(dict.fromkeys(
            aliases.get(name, name) for name in landscape.get("entities", [])
        ))
        landscape["entities"] = folded_entities
        answer_id = str(landscape["answer_id"])
        for position, entity in enumerate(folded_entities, 1):
            entity_answers[entity].append(answer_id)
            entity_positions[entity].append(position)
    minimum_recurrence = 2 if len(answers) > 1 else 1
    recurring_dimensions = [name for name, count in dimension_counts.most_common() if count >= minimum_recurrence]
    recurring_subquestions = [name for name, count in subquestion_counts.most_common() if count >= minimum_recurrence][:20]
    result = {
        "schema_version": "2.3.0",
        "job_id": job_id,
        "question": question,
        "captured_at": datetime.now(timezone.utc).isoformat(),
        "analysis_status": "available" if len(answers) >= 2 else "partial",
        "answers": answers,
        "source_observations": source_observations,
        "derived": {
            "answer_landscape": landscapes,
            "competitors": [
                {"name": name, "answer_ids": list(dict.fromkeys(answer_ids))}
                for name, answer_ids in sorted(entity_answers.items(), key=lambda item: (-len(item[1]), item[0]))
            ],
            "candidate_order_signals": [
                {"entity": name, "positions": positions}
                for name, positions in sorted(entity_positions.items(), key=lambda item: (sum(item[1]) / len(item[1]), item[0]))
            ],
            "recurring_dimensions": recurring_dimensions,
            "recurring_subquestions": recurring_subquestions,
            "unanswered_subquestions": list(dict.fromkeys(unanswered))[:20],
            "warnings": [] if source_observations else ["监控答案未提供可识别的公网引用链接；答案仍用于候选与子问题分析"],
        },
    }
    return result


def rows_to_answers(rows: list[tuple[str, dict[str, Any]]], target_question: str) -> list[dict[str, Any]]:
    discovered = Counter(
        normal(field(row, QUESTION_ALIASES)) for _platform, row in rows if normal(field(row, QUESTION_ALIASES))
    )
    target_key = question_key(target_question)
    exact_questions = {question for question in discovered if question_key(question) == target_key}
    if discovered and not exact_questions:
        choices = " | ".join(question for question, _count in discovered.most_common(10))
        raise ValueError(f"research_question_mismatch: no exact normalized monitoring question; found: {choices}")
    result: list[dict[str, Any]] = []
    for sheet_platform, row in rows:
        row_question = normal(field(row, QUESTION_ALIASES))
        if discovered and row_question not in exact_questions:
            continue
        platform = normal(field(row, PLATFORM_ALIASES)) or sheet_platform
        model = normal(field(row, MODEL_ALIASES)) or None
        sampled = field(row, DATE_ALIASES)
        for slot, value in answer_fields(row):
            text = normal(value)
            if text:
                result.append({
                    "answer_id": f"answer_{len(result) + 1:03d}",
                    "platform": platform,
                    "model": model or slot.split("_")[0] if slot else None,
                    "sampled_at": sampled,
                    "answer_text": text,
                    "citations": [],
                })
    if not result:
        raise ValueError("monitoring_answers_empty: no answer/content column contained readable text")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Normalize E2 AI monitoring answers")
    parser.add_argument("--content-job", type=Path, required=True)
    parser.add_argument("--monitoring-answers", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    job = json.loads(args.content_job.read_text(encoding="utf-8"))
    task = job.get("task") if isinstance(job.get("task"), dict) else {}
    question = normal(task.get("question_text"))
    if job.get("entry_category") == "foundation_start" or not question:
        raise ValueError("foundation_start skips E2; ordinary E2 requires a formal question")
    path = args.monitoring_answers
    if path.is_symlink() or not path.is_file():
        raise ValueError("monitoring answers must be a regular file")
    if path.suffix.casefold() == ".xlsx":
        rows, existing = rows_from_xlsx(path), None
    elif path.suffix.casefold() == ".csv":
        rows, existing = rows_from_csv(path), None
    elif path.suffix.casefold() == ".json":
        rows, existing = rows_from_json(path)
    else:
        raise ValueError("monitoring answers must be XLSX, CSV or JSON")
    if existing is not None:
        payload = normalize_existing(existing, str(job.get("job_id")), question)
    else:
        payload = build_context(str(job.get("job_id")), question, rows_to_answers(rows, question))
    validate_root(payload, "monitoring_context.schema.json", "E2 monitoring context")
    write_json(args.output, payload)
    print(json.dumps({
        "stage": "E2",
        "status": "completed" if payload["analysis_status"] == "available" else "completed_with_limits",
        "answer_count": len(payload["answers"]),
        "platform_count": len({item["platform"] for item in payload["answers"]}),
        "candidate_signal_count": len(payload["derived"]["competitors"]),
        "output": str(args.output),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
