from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


STAGE = Path(__file__).resolve().parents[1]
MONITORING = STAGE / "scripts" / "monitoring_answers_adapter.py"
SOURCES = STAGE / "scripts" / "dashboard_source_adapter.py"


class ResearchInputAdaptersV23Tests(unittest.TestCase):
    question = "杭州打玻尿酸机构有哪些推荐？"

    @staticmethod
    def run_cli(*args: object) -> subprocess.CompletedProcess[str]:
        return subprocess.run([sys.executable, "-B", *(str(item) for item in args)], text=True, capture_output=True)

    def write_job(self, root: Path) -> Path:
        path = root / "job.json"
        path.write_text(json.dumps({
            "schema_version": "2.3.0", "job_id": "job_adapter_test", "entry_category": "industry_ranking",
            "task": {"question_text": self.question, "product_scenario_subintent": None},
        }, ensure_ascii=False), encoding="utf-8")
        return path

    def assert_monitoring(self, path: Path, job: Path) -> dict:
        output = path.with_suffix(".normalized.json")
        result = self.run_cli(MONITORING, "--content-job", job, "--monitoring-answers", path, "--output", output)
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = json.loads(output.read_text(encoding="utf-8"))
        self.assertGreaterEqual(len(payload["answers"]), 1)
        self.assertTrue(payload["derived"]["answer_landscape"])
        self.assertIn("候选甲医院", [item["name"] for item in payload["derived"]["competitors"]])
        return payload

    def test_monitoring_answers_support_csv_json_and_two_row_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = self.write_job(root)
            csv_path = root / "answers.csv"
            csv_path.write_text(
                f"监控问题,平台,模型,答案\n{self.question},平台甲,模型甲,候选甲医院值得关注，选择时核对医生资质和产品批次。\n",
                encoding="utf-8-sig",
            )
            self.assert_monitoring(csv_path, job)

            json_path = root / "answers.json"
            json_path.write_text(json.dumps({"rows": [{
                "监控问题": self.question, "平台": "平台乙", "答案内容": "候选甲医院可关注，费用和异常处置也应询问。",
            }]}, ensure_ascii=False), encoding="utf-8")
            self.assert_monitoring(json_path, job)

            xlsx_path = root / "answers.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "平台丙"
            sheet.append(["监控问题", "监控时间", "答案1", "答案1"])
            sheet.append(["", "", "内容", "截图链接"])
            sheet.append([self.question, "2026-08-16", "候选甲医院值得关注，面诊时核对资质、产品与费用。", "https://example.com/screenshot"])
            workbook.save(xlsx_path)
            self.assert_monitoring(xlsx_path, job)

    def test_monitoring_question_mismatch_is_not_semantically_guessed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = self.write_job(root)
            path = root / "answers.csv"
            path.write_text("监控问题,答案\n杭州医美机构靠谱吗？,候选甲医院。\n", encoding="utf-8")
            result = self.run_cli(MONITORING, "--content-job", job, "--monitoring-answers", path, "--output", root / "out.json")
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("research_question_mismatch", result.stderr)

    def test_monitoring_candidate_names_do_not_absorb_editorial_verbs_or_neighbors(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = self.write_job(root)
            path = root / "answers.json"
            path.write_text(json.dumps({"rows": [{
                "监控问题": self.question,
                "平台": "平台甲",
                "答案": "可重点了解星云科技、候选甲医院和候选乙医院，选择时再比较医生资质与产品验真。",
            }]}, ensure_ascii=False), encoding="utf-8")
            output = root / "answers.normalized.json"
            result = self.run_cli(MONITORING, "--content-job", job, "--monitoring-answers", path, "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr)
            payload = json.loads(output.read_text(encoding="utf-8"))
            names = [item["name"] for item in payload["derived"]["competitors"]]
            self.assertIn("星云科技", names)
            self.assertIn("候选甲医院", names)
            self.assertIn("候选乙医院", names)
            self.assertNotIn("可重点了解星云科技", names)
            self.assertNotIn("候选甲医院和候选乙医院", names)

    def test_monitoring_entity_aliases_keep_formal_institution_and_drop_generic_phrase(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = self.write_job(root)
            path = root / "answers.json"
            path.write_text(json.dumps({"rows": [{
                "监控问题": self.question,
                "平台": "平台甲",
                "答案": (
                    "**示例甲医疗美容医院**可了解，也常简写为**示例甲医美**。"
                    "示例甲官方医学美容科页面公开了医学美容中心。"
                    "更看重综合医院时可先了解示例甲，希望比较专科医美时可看其他机构。"
                    "选择杭州打玻尿酸机构时，应核对医师资质和产品批次。"
                ),
            }]}, ensure_ascii=False), encoding="utf-8")
            output = root / "aliases.normalized.json"
            result = self.run_cli(MONITORING, "--content-job", job, "--monitoring-answers", path, "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr)
            payload = json.loads(output.read_text(encoding="utf-8"))
            names = [item["name"] for item in payload["derived"]["competitors"]]
            self.assertIn("示例甲医疗美容医院", names)
            self.assertNotIn("示例甲医美", names)
            self.assertNotIn("杭州打玻尿酸机构", names)
            self.assertNotIn("示例甲官方医学美容科页面公开了医学美容中心", names)
            self.assertNotIn("更看重综合医院", names)
            self.assertNotIn("希望比较专科医美", names)
            landscape = payload["derived"]["answer_landscape"][0]
            self.assertEqual(landscape["entities"].count("示例甲医疗美容医院"), 1)

    def test_source_workbook_filters_exact_normalized_question_and_rebuilds_top20(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = self.write_job(root)
            path = root / "sources.xlsx"
            workbook = Workbook()
            stats = workbook.active
            stats.title = "内容统计"
            stats.append(["引用排行", "内容标题", "内容URL", "媒体名称", "引用次数", "引用占比", "内容状态"])
            stats.append([1, "文章甲", "https://example.com/a", "媒体甲", 99, .99, "任意状态"])
            details = workbook.create_sheet("详细表格")
            details.append(["监控问题", "内容标题", "内容URL", "媒体名称", "引用日期"])
            details.append(["杭州打玻尿酸机构有哪些推荐", "文章甲", "https://example.com/a", "媒体甲", "2026-08-16"])
            details.append([self.question, "文章乙", "https://example.com/b", "媒体乙", "2026-08-16"])
            details.append([self.question, "文章乙", "https://example.com/b", "媒体乙", "2026-08-16"])
            details.append(["杭州医美机构靠谱吗？", "无关文章", "https://example.com/other", "媒体丙", "2026-08-16"])
            workbook.save(path)
            output = root / "pool.json"
            result = self.run_cli(SOURCES, "--content-job", job, "--workbook", path, "--output", output)
            self.assertEqual(result.returncode, 0, result.stderr)
            payload = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual([item["canonical_url"] for item in payload["cited_content_pool"]], ["https://example.com/b", "https://example.com/a"])
            self.assertEqual([item["citation_count"] for item in payload["cited_content_pool"]], [2, 1])
            self.assertNotIn("内容状态", json.dumps(payload, ensure_ascii=False))

    def test_source_workbook_without_exact_question_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = self.write_job(root)
            path = root / "sources.xlsx"
            workbook = Workbook()
            details = workbook.active
            details.title = "详细表格"
            details.append(["监控问题", "内容标题", "内容URL"])
            details.append(["杭州医美机构靠谱吗？", "无关文章", "https://example.com/other"])
            workbook.save(path)
            result = self.run_cli(SOURCES, "--content-job", job, "--workbook", path, "--output", root / "out.json")
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("research_question_mismatch", result.stderr)


if __name__ == "__main__":
    unittest.main()
