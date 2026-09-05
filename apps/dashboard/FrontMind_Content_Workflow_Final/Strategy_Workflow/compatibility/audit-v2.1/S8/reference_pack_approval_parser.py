#!/usr/bin/env python3
"""Parse and validate an S8 confirmation workbook into an approval receipt."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: install it in the workspace Python runtime") from exc


REVIEW_SHEETS = ["声明与证据", "受众与话语", "问题覆盖", "视觉资产"]
CONFIRMATION_LABELS = [
    "第一方资料公开授权与声明边界已确认",
    "自适应决策情境与适用范围已确认",
    "话语与竞品比较边界已确认",
    "七意图与五入口参考资产已确认",
    "图片权利与Logo规则已确认",
    "未决风险与附带条件已确认",
]
PACKABLE = {"approved"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def split_lines(value: Any) -> list[str]:
    raw = cell_text(value)
    return [line.strip(" -•\t") for line in raw.splitlines() if line.strip(" -•\t")]


def sheet_map(ws: Any) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        key = cell_text(row[0])
        if key:
            result[key] = row[1]
    return result


def parse_meta(ws: Any) -> dict[str, Any]:
    raw = sheet_map(ws)
    result: dict[str, Any] = {
        "schema_version": raw.get("schema_version"),
        "brand": raw.get("brand"),
        "pattern_registry_sha256": raw.get("pattern_registry_sha256"),
        "pattern_research_index_sha256": raw.get("pattern_research_index_sha256"),
    }
    try:
        result["input_files"] = json.loads(cell_text(raw.get("input_files")))
    except json.JSONDecodeError:
        result["input_files"] = []
    return result


def header_positions(ws: Any) -> tuple[int, dict[str, int]] | None:
    for row in range(1, min(ws.max_row, 12) + 1):
        values = {cell_text(ws.cell(row, column).value): column for column in range(1, ws.max_column + 1)}
        if "对象ID" in values and "审核决定" in values:
            return row, values
    return None


def parse_review_sheet(ws: Any) -> tuple[list[dict[str, Any]], list[str]]:
    parsed: list[dict[str, Any]] = []
    errors: list[str] = []
    found = header_positions(ws)
    if not found:
        return parsed, [f"{ws.title}: review headers not found"]
    header_row, columns = found
    id_col = columns["对象ID"]
    decision_col = columns["审核决定"]
    request_col = columns.get("修改要求")
    notes_col = columns.get("备注")
    for row in range(header_row + 1, ws.max_row + 1):
        object_id = cell_text(ws.cell(row, id_col).value)
        if not object_id:
            continue
        decision = cell_text(ws.cell(row, decision_col).value).lower()
        request = cell_text(ws.cell(row, request_col).value) if request_col else ""
        notes = cell_text(ws.cell(row, notes_col).value) if notes_col else ""
        if decision not in {"keep", "change", "block"}:
            errors.append(f"{ws.title}:{object_id} has invalid/blank decision")
        if decision in {"change", "block"} and not request:
            errors.append(f"{ws.title}:{object_id} needs a modification reason")
        parsed.append({"sheet": ws.title, "object_id": object_id, "decision": decision, "requested_change": request, "notes": notes})
    return parsed, errors


def main() -> int:
    parser = argparse.ArgumentParser(description="Parse S8 Reference Pack approval workbook")
    parser.add_argument("--confirmed", required=True)
    parser.add_argument("--brand")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    workbook_path = Path(args.confirmed).expanduser().resolve()
    if not workbook_path.is_file():
        parser.error(f"workbook does not exist: {workbook_path}")
    try:
        workbook = load_workbook(workbook_path, data_only=True)
    except Exception as exc:
        parser.error(f"cannot open workbook: {exc}")

    required_sheets = set(REVIEW_SHEETS + ["审批决议", "_meta"])
    missing = required_sheets - set(workbook.sheetnames)
    if missing:
        parser.error("missing sheets: " + ", ".join(sorted(missing)))

    meta = parse_meta(workbook["_meta"])
    decision_values = sheet_map(workbook["审批决议"])
    brand = args.brand or cell_text(decision_values.get("品牌")) or cell_text(meta.get("brand"))
    status = cell_text(decision_values.get("审批状态")).lower()
    reviewer = cell_text(decision_values.get("负责人"))
    reviewer_role = cell_text(decision_values.get("负责人角色"))
    approved_at = cell_text(decision_values.get("审批日期"))
    conditions = split_lines(decision_values.get("附带条件（每行一项）"))
    risks = split_lines(decision_values.get("未决风险（每行一项）"))
    confirmations = {label: cell_text(decision_values.get(label)).upper() == "YES" for label in CONFIRMATION_LABELS}

    errors: list[str] = []
    if status not in {"approved", "changes_requested"}:
        errors.append(f"invalid approval status: {status or 'blank'}")
    if not brand:
        errors.append("brand is required")
    if status in PACKABLE and (not reviewer or not reviewer_role or not approved_at):
        errors.append("packable approval needs reviewer, reviewer role, and approval date")
    if status in PACKABLE and not all(confirmations.values()):
        errors.append("all six confirmations must be YES for a packable approval")
    if not isinstance(meta.get("input_files"), list) or not meta["input_files"]:
        errors.append("_meta input file fingerprints are missing")
    for field in ("pattern_registry_sha256", "pattern_research_index_sha256"):
        value = cell_text(meta.get(field)).lower()
        if len(value) != 64 or any(character not in "0123456789abcdef" for character in value):
            errors.append(f"_meta {field} is missing or invalid")

    reviews: list[dict[str, Any]] = []
    for sheet_name in REVIEW_SHEETS:
        rows, sheet_errors = parse_review_sheet(workbook[sheet_name])
        reviews.extend(rows)
        errors.extend(sheet_errors)
    pending = [row for row in reviews if row["decision"] in {"change", "block"}]
    if status in PACKABLE and pending:
        errors.append(f"packable approval has {len(pending)} change/block decisions")
    if status == "approved" and (conditions or risks):
        errors.append("approved cannot contain conditions or unresolved risks; use changes_requested")

    if errors:
        print(json.dumps({"valid": False, "errors": errors, "status": status, "pending_decisions": len(pending)}, ensure_ascii=False, indent=2))
        return 1

    receipt = {
        "schema_version": "2.0.0",
        "brand": brand,
        "approval": {
            "status": status,
            "reviewer": reviewer,
            "reviewer_role": reviewer_role,
            "approved_at": approved_at,
            "conditions": conditions,
            "unresolved_risks": risks,
            "notes": cell_text(decision_values.get("总备注")),
        },
        "confirmations": confirmations,
        "input_files": meta["input_files"],
        "pattern_registry_sha256": meta.get("pattern_registry_sha256"),
        "pattern_research_index_sha256": meta.get("pattern_research_index_sha256"),
        "review_decisions": reviews,
        "workbook": {"file_name": workbook_path.name, "sha256": sha256(workbook_path)},
    }
    output = Path(args.out).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8") as handle:
        json.dump(receipt, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(json.dumps({"valid": True, "status": status, "output": str(output), "reviewed_objects": len(reviews)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
