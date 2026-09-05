#!/usr/bin/env python3
"""Generate the human review workbook for an S8 Reference Pack approval."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required: install it in the workspace Python runtime") from exc


NAVY = "183153"
BLUE = "2E75B6"
LIGHT_BLUE = "D9EAF7"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"


ROLE_PATTERNS = {
    "adapter_report": ("**/S2_*adapter_report.json",),
    "knowledge_registry": ("**/S2_*knowledge_registry.json",),
    "source_registry": ("**/S2_*source_registry.json",),
    "claim_registry": ("**/S2_*claim_registry.json",),
    "image_registry": ("**/S2_*image_registry.json",),
    "brand_facts": ("**/S4_*brand_facts.json", "**/S4_*canonical_brand_facts*.json"),
    "information_evidence_architecture": ("**/S4_*information_evidence*.json", "**/S4_*信息与证据架构*.json"),
    "voice_contract": ("**/S5_*voice_contract*.json", "**/S5_*文章话语契约*.json"),
    "visual_reference": ("**/S6_*visual_reference*.json", "**/S6_*视觉参考体系*.json"),
    "question_library": ("**/S7_*question_library*.json", "**/S7_*问题与FAQ参考库*.json"),
}

OPTIONAL_ROLE_PATTERNS = {
    "trend_reference": ("**/S3_*trend_reference*.json", "**/S3_*内容趋势参考*.json", "**/S3_status.json", "**/S3_*status.json"),
}


def load(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def find_one(root: Path, patterns: tuple[str, ...]) -> Path | None:
    matches: set[Path] = set()
    for pattern in patterns:
        matches.update(path.resolve() for path in root.glob(pattern) if path.is_file())
    return sorted(matches, key=lambda path: (path.stat().st_mtime_ns, str(path)), reverse=True)[0] if matches else None


def text(value: Any, limit: int = 12000) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        result = value
    elif isinstance(value, (list, dict)):
        result = json.dumps(value, ensure_ascii=False)
    else:
        result = str(value)
    return result if len(result) <= limit else result[:limit] + "…"


def item_list(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, dict):
        for key in ("knowledge_units", "sources", "claims", "images", "items"):
            if isinstance(payload.get(key), list):
                return [item for item in payload[key] if isinstance(item, dict)]
    return []


def list_values(payload: Any, key: str) -> set[str]:
    result: set[str] = set()
    if isinstance(payload, dict):
        for name, child in payload.items():
            if name == key and isinstance(child, list):
                result.update(str(value) for value in child if isinstance(value, str))
            result.update(list_values(child, key))
    elif isinstance(payload, list):
        for child in payload:
            result.update(list_values(child, key))
    return result


def scalar_values(payload: Any, key: str) -> set[str]:
    result: set[str] = set()
    if isinstance(payload, dict):
        for name, child in payload.items():
            if name == key and isinstance(child, str):
                result.add(child)
            result.update(scalar_values(child, key))
    elif isinstance(payload, list):
        for child in payload:
            result.update(scalar_values(child, key))
    return result


def workbook_base() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "FrontMind S8 Reference Pack"
    workbook.properties.title = "Reference Pack 确认表"
    return workbook


def title(ws: Any, title_text: str, subtitle: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = ws.cell(1, 1, title_text)
    cell.font = Font(size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws.cell(2, 1, subtitle).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.row_dimensions[2].height = 34


def header(ws: Any, row: int, headers: list[str]) -> None:
    for column, value in enumerate(headers, 1):
        cell = ws.cell(row, column, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = f"A{row}:{ws.cell(row, len(headers)).coordinate}"
    ws.freeze_panes = f"A{row + 1}"


def finish_sheet(ws: Any, widths: list[int], review_columns: tuple[int, int, int] | None = None) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if review_columns:
        decision_col, request_col, notes_col = review_columns
        for row in range(4, ws.max_row + 1):
            for column in review_columns:
                ws.cell(row, column).fill = PatternFill("solid", fgColor=YELLOW)
        validation = DataValidation(type="list", formula1='"keep,change,block"', allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f"{ws.cell(4, decision_col).coordinate}:{ws.cell(max(4, ws.max_row), decision_col).coordinate}")
        ws.conditional_formatting.add(
            f"{ws.cell(4, decision_col).coordinate}:{ws.cell(max(4, ws.max_row), notes_col).coordinate}",
            FormulaRule(formula=[f'${ws.cell(4, decision_col).column_letter}4="block"'], fill=PatternFill("solid", fgColor=RED)),
        )


def build_overview(workbook: Workbook, brand: str, files: dict[str, Path], payloads: dict[str, Any], gates: list[str]) -> None:
    ws = workbook.create_sheet("说明与总览")
    title(ws, f"{brand} Reference Pack 确认", "只确认将被文章消费的事实、证据、受众、话语、问题与视觉边界。黄色单元格由审核人填写。", 5)
    header(ws, 3, ["角色", "文件", "SHA-256", "对象数", "状态"])
    for role, path in files.items():
        payload = payloads[role]
        count = len(item_list(payload)) or (len(payload.get("questions", [])) if isinstance(payload, dict) else "")
        status = payload.get("status") or payload.get("mode") or "loaded" if isinstance(payload, dict) else "loaded"
        ws.append([role, path.name, sha256(path), count, status])
    row = ws.max_row + 2
    ws.cell(row, 1, "硬门槛")
    ws.cell(row, 1).font = Font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    if gates:
        for issue in gates:
            ws.append(["FAIL", issue, "", "", "必须回退修复"])
    else:
        ws.append(["PASS", "所有自动硬门槛通过；仍需负责人逐项确认", "", "", "可进入人审"])
    finish_sheet(ws, [24, 58, 68, 12, 22])


def build_claims(
    workbook: Workbook,
    claims: Any,
    architecture: Any,
    brand_facts: Any,
    sources: Any,
) -> None:
    ws = workbook.create_sheet("声明与证据")
    headers = ["对象ID", "对象类型", "声明/信息", "发布状态", "证据精度", "来源/Proof refs", "允许措辞/原因", "审核决定", "修改要求", "备注"]
    title(ws, "声明与证据", "keep=保留；change=回 S2/S4 修改；block=禁止进入文章。", len(headers))
    header(ws, 3, headers)
    for source in item_list(sources):
        if source.get("source_origin_class") not in {"first_party_official", "first_party_internal"}:
            continue
        ws.append([
            source.get("source_id"),
            "source_authorization",
            text({"title": source.get("title"), "publisher": source.get("publisher")}),
            source.get("publication_authorization"),
            source.get("evidence_precision"),
            text({
                "source_type": source.get("source_type"),
                "source_origin_class": source.get("source_origin_class"),
                "url": source.get("url"),
                "file_path": source.get("file_path"),
                "locator": source.get("locator"),
            }),
            text({"authorization_audit": source.get("notes"), "content_status": source.get("content_status")}),
            "", "", "",
        ])
    for item in item_list(claims):
        ws.append([
            item.get("claim_id"), "claim", text(item.get("claim_text"), 4000), item.get("allowed_usage"), item.get("evidence_precision"),
            text(item.get("source_ids") or item.get("evidence_document_refs")), text(item.get("reason")), "", "", "",
        ])
    for proof in architecture.get("proof_bundles", []) if isinstance(architecture, dict) else []:
        if not isinstance(proof, dict):
            continue
        ws.append([
            proof.get("proof_id"), "proof_bundle", text(proof.get("statement")), proof.get("allowed_usage"), proof.get("evidence_precision"),
            text({"claims": proof.get("claim_ids"), "sources": proof.get("source_ids")}), text({"wording": proof.get("allowed_wording"), "caveat": proof.get("caveat")}), "", "", "",
        ])
    if isinstance(brand_facts, dict):
        identity = brand_facts.get("brand_identity") or {}
        positioning = brand_facts.get("positioning") or {}
        ws.append(["brand:identity", "brand_fact", text(identity), "canonical", "source", text(identity.get("source_ids")), "品牌身份", "", "", ""])
        ws.append(["brand:positioning", "brand_fact", text(positioning), "canonical", "source", text(positioning.get("source_ids")), "定位与受众", "", "", ""])
        for section, id_key in (("offerings", "offering_id"), ("capabilities", "capability_id"), ("proof_points", "proof_id"), ("limitations", "limitation_id")):
            for item in brand_facts.get(section, []):
                if isinstance(item, dict):
                    ws.append([item.get(id_key), f"brand_fact:{section}", text(item), "canonical", "source", text({"sources": item.get("source_ids"), "claims": item.get("claim_ids")}), "S4 canonical brand facts", "", "", ""])
    finish_sheet(ws, [24, 16, 56, 22, 16, 48, 50, 14, 44, 32], (8, 9, 10))


def build_audience_voice(workbook: Workbook, architecture: Any, voice: Any) -> None:
    ws = workbook.create_sheet("受众与话语")
    headers = ["对象ID", "对象类型", "内容", "上游refs", "边界/规则", "审核决定", "修改要求", "备注"]
    title(ws, "受众与话语", "确认文章写给谁、决策标准和怎样表达；修改分别回 S4 或 S5。", len(headers))
    header(ws, 3, headers)
    if isinstance(architecture, dict):
        for context in architecture.get("audience_decision_contexts", []):
            if isinstance(context, dict):
                ws.append([context.get("context_id"), "decision_context", text({key: context.get(key) for key in ("role", "trigger", "job", "constraints", "decision_criteria")}), text({"knowledge_ids": context.get("knowledge_ids"), "source_ids": context.get("source_ids")}), context.get("provenance_status"), "", "", ""])
        for pillar in architecture.get("message_pillars", []):
            if isinstance(pillar, dict):
                ws.append([pillar.get("message_id") or pillar.get("pillar_id"), "message_pillar", text(pillar.get("message")), text(pillar.get("proof_refs")), text(pillar.get("must_not_claim") or pillar.get("boundary")), "", "", ""])
        value_structure = architecture.get("value_structure") or {}
        for layer in ("functional", "emotional", "self_expression"):
            for value in value_structure.get(layer, []) if isinstance(value_structure, dict) else []:
                if isinstance(value, dict):
                    ws.append([
                        value.get("value_id"), f"value:{layer}", text(value.get("statement")),
                        text({"contexts": value.get("audience_context_refs"), "proofs": value.get("proof_refs")}),
                        text({"mode": value.get("expression_mode"), "allowed": value.get("allowed_wording"), "must_not": value.get("must_not_imply")}),
                        "", "", "",
                    ])
        for item in architecture.get("differentiation_evidence_map", []):
            if isinstance(item, dict):
                ws.append([
                    item.get("differentiation_id"), "differentiation", text({"dimension": item.get("dimension"), "statement": item.get("statement"), "against": item.get("compared_against")}),
                    text({"contexts": item.get("audience_context_refs"), "proofs": item.get("proof_refs")}),
                    text({"state": item.get("evidence_state"), "allowed": item.get("allowed_wording"), "must_not": item.get("must_not_claim"), "as_of": item.get("as_of")}),
                    "", "", "",
                ])
        for angle in architecture.get("brand_priority_angles", []):
            if isinstance(angle, dict):
                ws.append([
                    angle.get("angle_id"), "brand_priority_angle", text(angle.get("angle")),
                    text({"contexts": angle.get("audience_context_refs"), "values": angle.get("value_refs"), "differentiation": angle.get("differentiation_refs"), "proofs": angle.get("proof_refs")}),
                    text({"entries": angle.get("eligible_entry_categories"), "allowed": angle.get("allowed_wording"), "must_not": angle.get("must_not_claim")}),
                    "", "", "",
                ])
    if isinstance(voice, dict):
        ws.append(["voice:default", "voice_contract", text(voice.get("default_voice")), text(voice.get("upstream_refs")), "默认规则", "", "", ""])
        ws.append(["voice:terminology", "terminology", text(voice.get("terminology")), "", "术语与禁用词", "", "", ""])
        ws.append(["voice:comparison", "comparison_rules", text(voice.get("comparison_rules")), "", "竞品比较边界", "", "", ""])
    finish_sheet(ws, [24, 20, 68, 40, 38, 14, 44, 32], (6, 7, 8))


def build_questions(workbook: Workbook, library: Any, pattern_registry: Any) -> None:
    ws = workbook.create_sheet("问题覆盖")
    headers = ["对象ID", "问题", "主意图", "入口", "子意图", "候选Pattern", "Provenance", "可回答性", "Answer thesis（非完整答案）", "审核决定", "修改要求", "备注"]
    title(ws, "问题覆盖", "必须 30–50 题，覆盖共享 registry 七意图与五入口参考资产；每题有 provenance。", len(headers))
    header(ws, 3, headers)
    for question in library.get("questions", []) if isinstance(library, dict) else []:
        if not isinstance(question, dict):
            continue
        boundary = question.get("answer_boundary") or {}
        ws.append([
            question.get("question_id"), question.get("canonical_question"), question.get("primary_intent_tag"), question.get("entry_category"),
            question.get("product_scenario_subintent_ref"), text(question.get("pattern_refs")), text(question.get("provenance")), question.get("answerability"),
            text(boundary.get("answer_thesis"), 500), "", "", "",
        ])
    coverage_row = ws.max_row + 2
    tags = [item.get("id") for item in pattern_registry.get("question_intent_tags", []) if isinstance(item, dict)] if isinstance(pattern_registry, dict) else []
    categories = [
        item.get("id")
        for item in pattern_registry.get("canonical_content_entries", [])
        if isinstance(item, dict)
    ] if isinstance(pattern_registry, dict) else []
    ws.cell(coverage_row, 1, "共享七意图")
    ws.cell(coverage_row, 2, text(tags))
    ws.cell(coverage_row + 1, 1, "共享五入口")
    ws.cell(coverage_row + 1, 2, text(categories))
    finish_sheet(ws, [22, 50, 18, 22, 24, 24, 50, 24, 50, 14, 44, 30], (10, 11, 12))


def build_visual(workbook: Workbook, visual: Any, assets: Any) -> None:
    ws = workbook.create_sheet("视觉资产")
    headers = ["对象ID", "来源/路径", "权利状态", "S6用法", "允许槽位", "说明", "审核决定", "修改要求", "备注"]
    title(ws, "视觉资产", "direct_use/overlay_only 只允许 approved 或 approved_with_credit；Logo 必须使用真实文件覆盖。", len(headers))
    header(ws, 3, headers)
    rights = {str(item.get("asset_id")): item for item in item_list(assets) if item.get("asset_id")}
    usage = visual.get("asset_usage", []) if isinstance(visual, dict) else []
    used = set()
    for item in usage:
        if not isinstance(item, dict):
            continue
        asset_id = str(item.get("asset_ref") or "")
        source = rights.get(asset_id, {})
        used.add(asset_id)
        ws.append([asset_id, source.get("file_path") or source.get("url"), source.get("rights_status"), item.get("usage_mode"), text(item.get("allowed_slots")), item.get("rationale"), "", "", ""])
    for asset_id, source in rights.items():
        if asset_id not in used:
            ws.append([asset_id, source.get("file_path") or source.get("url"), source.get("rights_status"), "not_selected", "", source.get("caption"), "", "", ""])
    logo_policy = visual.get("logo_policy") if isinstance(visual, dict) else None
    ws.append(["visual:logo-policy", "", "", "overlay_real_asset", "logo slots", text(logo_policy), "", "", ""])
    finish_sheet(ws, [26, 52, 22, 22, 34, 52, 14, 44, 30], (7, 8, 9))


def build_approval(workbook: Workbook, brand: str) -> None:
    ws = workbook.create_sheet("审批决议")
    title(ws, "最终审批决议", "只有六项全部 YES 且无 change/block 时才可 approved。黄色单元格由负责人填写。", 4)
    rows = [
        ("品牌", brand),
        ("第一方资料公开授权与声明边界已确认", "NO"),
        ("自适应决策情境与适用范围已确认", "NO"),
        ("话语与竞品比较边界已确认", "NO"),
        ("七意图与五入口参考资产已确认", "NO"),
        ("图片权利与Logo规则已确认", "NO"),
        ("未决风险与附带条件已确认", "NO"),
        ("审批状态", "changes_requested"),
        ("负责人", ""),
        ("负责人角色", ""),
        ("审批日期", date.today().isoformat()),
        ("附带条件（每行一项）", ""),
        ("未决风险（每行一项）", ""),
        ("总备注", ""),
    ]
    header(ws, 3, ["字段", "填写值", "说明", "机器规则"])
    for label, default in rows:
        ws.append([label, default, "", ""])
    for row in range(4, ws.max_row + 1):
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=YELLOW)
    yes_no = DataValidation(type="list", formula1='"YES,NO"')
    ws.add_data_validation(yes_no)
    yes_no.add("B5:B10")
    status = DataValidation(type="list", formula1='"approved,changes_requested"')
    ws.add_data_validation(status)
    status.add("B11")
    ws.freeze_panes = "A4"
    finish_sheet(ws, [38, 56, 44, 38])


def gate_checks(payloads: dict[str, Any], pattern_registry: Any) -> list[str]:
    issues: list[str] = []
    report = payloads["adapter_report"]
    if report.get("fatal_errors"):
        issues.append("S2 adapter_report contains fatal_errors")
    if report.get("formal_handoff_ready") is not True:
        issues.append("S2 adapter_report is not formal_handoff_ready")
    knowledge_ids = {item.get("knowledge_id") for item in item_list(payloads["knowledge_registry"]) if item.get("knowledge_id")}
    source_ids = {item.get("source_id") for item in item_list(payloads["source_registry"]) if item.get("source_id")}
    claim_ids = {item.get("claim_id") for item in item_list(payloads["claim_registry"]) if item.get("claim_id")}
    source_records = {
        item.get("source_id"): item for item in item_list(payloads["source_registry"]) if item.get("source_id")
    }
    claim_sources = {
        item.get("claim_id"): set(item.get("source_ids") or [])
        for item in item_list(payloads["claim_registry"]) if item.get("claim_id")
    }
    allowed_origin_classes = {
        "first_party_official", "first_party_internal", "third_party_public", "runtime_approved",
    }
    for source_id, source in source_records.items():
        origin_class = source.get("source_origin_class")
        authorization = source.get("publication_authorization")
        if origin_class not in allowed_origin_classes:
            issues.append(f"S2 source {source_id} has invalid source_origin_class: {origin_class!r}")
        if origin_class == "first_party_official" and authorization not in {"public_approved", "anonymized_approved"}:
            issues.append(f"S2 source {source_id} has invalid first-party official authorization: {authorization!r}")
        if origin_class == "first_party_internal" and authorization != "internal_only":
            issues.append(f"S2 source {source_id} has invalid first-party internal authorization: {authorization!r}")
        if origin_class == "runtime_approved" and authorization != "runtime_approved":
            issues.append(f"S2 source {source_id} has invalid runtime authorization: {authorization!r}")
    image_ids = {item.get("asset_id") for item in item_list(payloads["image_registry"]) if item.get("asset_id")}
    pattern_ids = {item.get("id") for item in pattern_registry.get("patterns", []) if isinstance(item, dict) and item.get("id")}
    architecture = payloads["information_evidence_architecture"]
    proof_ids = {item.get("proof_id") for item in architecture.get("proof_bundles", []) if isinstance(item, dict)}
    context_ids = {item.get("context_id") for item in architecture.get("audience_decision_contexts", []) if isinstance(item, dict)}
    for index, pillar in enumerate(architecture.get("message_pillars", [])):
        if not isinstance(pillar, dict) or not set(pillar.get("proof_refs") or []) & proof_ids:
            issues.append(f"S4 message_pillars[{index}] has no valid proof")
    if not architecture.get("brand_priority_angles"):
        issues.append("S4 brand_priority_angles is empty")
    if not architecture.get("differentiation_evidence_map"):
        issues.append("S4 differentiation_evidence_map is empty")
    for index, proof in enumerate(architecture.get("proof_bundles", [])):
        if not isinstance(proof, dict):
            continue
        attached = set().union(*(claim_sources.get(claim_id, set()) for claim_id in proof.get("claim_ids", []))) if proof.get("claim_ids") else set()
        unattached = set(proof.get("source_ids") or []) - attached
        if unattached:
            issues.append(f"S4 proof_bundles[{index}] has sources not attached to claims: {sorted(unattached)}")
        if proof.get("allowed_usage") in {"public_fact", "public_with_qualification"}:
            unauthorized = sorted(
                source_id for source_id in proof.get("source_ids", [])
                if (source_records.get(source_id) or {}).get("source_type") in {"first_party_official", "first_party_internal"}
                and (source_records.get(source_id) or {}).get("publication_authorization") not in {
                    "public_approved", "anonymized_approved", "runtime_approved", "not_applicable",
                }
            )
            if unauthorized:
                issues.append(f"S4 proof_bundles[{index}] has non-publishable first-party sources: {unauthorized}")
    for label, payload in (("S4", architecture), ("S4 brand facts", payloads["brand_facts"])):
        unknown_sources = list_values(payload, "source_ids") - source_ids
        unknown_claims = list_values(payload, "claim_ids") - claim_ids
        if unknown_sources:
            issues.append(f"{label} references unknown source_ids: {sorted(unknown_sources)}")
        if unknown_claims:
            issues.append(f"{label} references unknown claim_ids: {sorted(unknown_claims)}")
    unknown_knowledge = list_values(architecture, "knowledge_ids") - knowledge_ids
    if unknown_knowledge:
        issues.append(f"S4 references unknown knowledge_ids: {sorted(unknown_knowledge)}")
    if list_values(architecture, "pattern_refs") - pattern_ids:
        issues.append("S4 references unknown pattern_ids")
    brand_facts = payloads["brand_facts"]
    for field in ("brand_identity", "positioning", "offerings", "capabilities", "proof_points", "limitations"):
        if field not in brand_facts:
            issues.append(f"S4 brand facts missing {field}")
    voice = payloads["voice_contract"]
    if set((voice.get("claim_expression_rules") or {}).keys()) < {"public_fact", "public_with_qualification", "internal_only", "do_not_use"}:
        issues.append("S5 lacks four evidence-state expression rules")
    if list_values(voice, "proof_refs") - proof_ids:
        issues.append("S5 references unknown S4 proof IDs")
    if list_values(voice, "pattern_refs") - pattern_ids:
        issues.append("S5 references unknown pattern IDs")
    rights = {item.get("asset_id"): item.get("rights_status") for item in item_list(payloads["image_registry"])}
    visual = payloads["visual_reference"]
    for item in visual.get("asset_usage", []):
        if isinstance(item, dict) and item.get("usage_mode") in {"direct_use", "overlay_only"} and rights.get(item.get("asset_ref")) not in {"approved", "approved_with_credit"}:
            issues.append(f"S6 direct asset lacks publishable rights: {item.get('asset_ref')}")
    if (visual.get("logo_policy") or {}).get("mode") != "overlay_real_asset":
        issues.append("S6 logo policy is not overlay_real_asset")
    visual_asset_refs = scalar_values(visual, "asset_ref") | set((visual.get("brand_constraints") or {}).get("logo_asset_refs") or [])
    if visual_asset_refs - image_ids:
        issues.append(f"S6 references unknown image IDs: {sorted(visual_asset_refs - image_ids)}")
    if list_values(visual, "pattern_refs") - pattern_ids:
        issues.append("S6 references unknown pattern IDs")
    library = payloads["question_library"]
    questions = library.get("questions", [])
    if not 30 <= len(questions) <= 50:
        issues.append(f"S7 question count is {len(questions)}, expected 30..50")
    required_tags = {item.get("id") for item in pattern_registry.get("question_intent_tags", []) if isinstance(item, dict)}
    required_categories = {
        item.get("id")
        for item in pattern_registry.get("canonical_content_entries", [])
        if isinstance(item, dict)
    }
    used_tags = {question.get("primary_intent_tag") for question in questions if isinstance(question, dict)}
    used_categories = {question.get("entry_category") for question in questions if isinstance(question, dict)}
    if required_tags - used_tags:
        issues.append("S7 does not cover all seven primary intent tags")
    if required_categories - used_categories:
        issues.append("S7 does not cover all five content entries")
    if any(not question.get("provenance") for question in questions if isinstance(question, dict)):
        issues.append("S7 has questions without provenance")
    if list_values(library, "source_ids") - source_ids:
        issues.append("S7 references unknown source IDs")
    if list_values(library, "claim_ids") - claim_ids:
        issues.append("S7 references unknown claim IDs")
    if list_values(library, "knowledge_ids") - knowledge_ids:
        issues.append("S7 references unknown knowledge IDs")
    if list_values(library, "pattern_refs") - pattern_ids:
        issues.append("S7 references unknown pattern IDs")
    if list_values(library, "audience_context_refs") - context_ids:
        issues.append("S7 references unknown S4 audience contexts")
    if list_values(library, "proof_refs") - proof_ids:
        issues.append("S7 references unknown S4 proof bundles")
    trend = payloads.get("trend_reference")
    if isinstance(trend, dict):
        if trend.get("status") not in {"completed", "skipped_not_needed", "skipped_no_reliable_signal"}:
            issues.append("S3 trend reference has invalid status")
        if list_values(trend, "source_ids") - source_ids or list_values(trend, "knowledge_ids") - knowledge_ids:
            issues.append("S3 trend reference has unknown registry IDs")
        if list_values(trend, "pattern_refs") - pattern_ids:
            issues.append("S3 trend reference has unknown pattern IDs")
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate S8 Reference Pack confirmation workbook")
    parser.add_argument("--brand", required=True)
    parser.add_argument("--work-dir", required=True)
    parser.add_argument("--pattern-registry", required=True)
    parser.add_argument("--pattern-research-index", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    root = Path(args.work_dir).expanduser().resolve()
    if not root.is_dir():
        parser.error(f"work directory does not exist: {root}")
    files: dict[str, Path] = {}
    for role, patterns in ROLE_PATTERNS.items():
        path = find_one(root, patterns)
        if not path:
            parser.error(f"missing required artifact: {role}")
        files[role] = path
    for role, patterns in OPTIONAL_ROLE_PATTERNS.items():
        path = find_one(root, patterns)
        if path:
            files[role] = path
    pattern_path = Path(args.pattern_registry).expanduser().resolve()
    if not pattern_path.is_file():
        parser.error(f"pattern registry does not exist: {pattern_path}")
    pattern_research_path = Path(args.pattern_research_index).expanduser().resolve()
    if not pattern_research_path.is_file():
        parser.error(f"pattern research index does not exist: {pattern_research_path}")
    files["pattern_registry"] = pattern_path
    files["pattern_research_index"] = pattern_research_path
    try:
        payloads = {role: load(path) for role, path in files.items()}
    except (OSError, json.JSONDecodeError) as exc:
        parser.error(f"invalid input JSON: {exc}")
    gates = gate_checks(payloads, payloads["pattern_registry"])
    if gates:
        parser.error("hard gates failed: " + "; ".join(gates))

    workbook = workbook_base()
    build_overview(workbook, args.brand, files, payloads, gates)
    build_claims(
        workbook,
        payloads["claim_registry"],
        payloads["information_evidence_architecture"],
        payloads["brand_facts"],
        payloads["source_registry"],
    )
    build_audience_voice(workbook, payloads["information_evidence_architecture"], payloads["voice_contract"])
    build_questions(workbook, payloads["question_library"], payloads["pattern_registry"])
    build_visual(workbook, payloads["visual_reference"], payloads["image_registry"])
    build_approval(workbook, args.brand)
    meta = workbook.create_sheet("_meta")
    meta.append(["key", "value"])
    meta.append(["schema_version", "2.0.0"])
    meta.append(["brand", args.brand])
    meta.append(["input_files", json.dumps([{"role": role, "file_name": path.name, "sha256": sha256(path), "bytes": path.stat().st_size} for role, path in files.items()], ensure_ascii=False)])
    meta.append(["pattern_registry_sha256", sha256(pattern_path)])
    meta.append(["pattern_research_index_sha256", sha256(pattern_research_path)])
    meta.sheet_state = "hidden"

    output = Path(args.out).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(json.dumps({"status": "ok", "output": str(output), "sheets": workbook.sheetnames, "inputs": len(files)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
